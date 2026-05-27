#!/usr/bin/env python3
"""תמחור פאנלים בעדיפות נוחות עבודה (הכי מעט חתיכות) + 30% רווח.

בחירת הפלטה לכל פאנל: הכי מעט חתיכות (תפרים) → הכי מעט לוחות → מחיר.
מציג גם כמה לוחות מכל סוג צריך לקנות, ואת ההפרש מול תכנון זול-ביותר.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from sheet_calc import PANELS, PRICES, split_panel
from build_cut_plan import best_plate, plates_for, compute_plan

MARKUP = 1.30
OBJ = "convenience"

rows, total_cost, total_sell = [], 0, 0
for (w, h, q) in PANELS:
    plate = best_plate(w, h, q, OBJ)
    pieces = len(split_panel(w, h, *plate))
    n = plates_for(w, h, q, plate)
    cost = n * PRICES[plate]
    unit_cost = cost / q
    unit_sell = round(unit_cost * MARKUP)
    line = unit_sell * q
    total_cost += cost
    total_sell += line
    rows.append((f"{w}x{h}", q, f"{plate[0]}x{plate[1]}", pieces, n,
                 cost, round(unit_cost, 1), unit_sell, line))

# לוחות לקנייה (אריזה משותפת) בשני שיקולים
conv_mix = compute_plan(OBJ)[2]
cost_mix = compute_plan("cost")[2]
conv_buy = sum(c * PRICES[s] for s, c in conv_mix.items())
cost_buy = sum(c * PRICES[s] for s, c in cost_mix.items())

print(f"תמחור ({OBJ}, +30%):")
for r in rows:
    print(f"  {r[0]:<12} x{r[1]:<3} {r[2]:<10} {r[3]} חת׳  {r[4]:>3} לוחות  "
          f"יח׳ {r[7]:>5}₪  סה״כ {r[8]:>6}₪")
print(f"\nלוחות לקנייה (נוחות): {sum(conv_mix.values())} | עלות {conv_buy:,}₪")
for s, c in sorted(conv_mix.items()):
    print(f"  {s[0]}x{s[1]}: {c}")
print(f"לעומת תכנון זול-ביותר: {sum(cost_mix.values())} לוחות | עלות {cost_buy:,}₪ "
      f"(הפרש {conv_buy-cost_buy:+,}₪)")
print(f"סה״כ הצעת מחיר: {total_sell:,}₪")

# ---------------------------------------------------------------- אקסל
navy = PatternFill("solid", fgColor="1F4E78")
gray = PatternFill("solid", fgColor="D9E1F2")
green = PatternFill("solid", fgColor="C6EFCE")
yellow = PatternFill("solid", fgColor="FFF2CC")
wbf = Font(bold=True, color="FFFFFF", size=11)
bold = Font(bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")
border = Border(*[Side(style="thin", color="B0B0B0")] * 4)
NIS = '#,##0" ₪"'

wb = Workbook()
ws = wb.active
ws.title = "תמחור"
ws.sheet_view.rightToLeft = True
heads = ["פאנל (מ״מ)", "כמות", "פלטה", "חתיכות לפאנל", "לוחות",
         "עלות חומר", "עלות ליחידה", "מחיר מכירה ליחידה", "סה״כ מכירה"]
widths = [16, 8, 12, 12, 8, 12, 12, 16, 14]
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
t = ws.cell(1, 1, "הצעת מחיר - עדיפות נוחות עבודה (כולל 30% רווח)")
t.fill = navy; t.font = Font(bold=True, color="FFFFFF", size=14); t.alignment = center
for c, (lab, w) in enumerate(zip(heads, widths), start=1):
    cell = ws.cell(2, c, lab)
    cell.fill = navy; cell.font = wbf; cell.alignment = center; cell.border = border
    ws.column_dimensions[L(c)].width = w
r = 3
for (name, q, plate, pieces, n, cost, uc, us, line) in rows:
    for c, v in enumerate([name, q, plate, pieces, n, cost, uc, us, line], start=1):
        cell = ws.cell(r, c, v); cell.alignment = center; cell.border = border
        if c in (6, 7):
            cell.number_format = NIS
        if c == 8:
            cell.number_format = NIS; cell.font = bold; cell.fill = yellow
        if c == 9:
            cell.number_format = NIS; cell.font = bold
    r += 1
ws.cell(r, 1, "סה״כ").font = bold
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.cell(r, 1).alignment = center; ws.cell(r, 1).fill = gray
tc = ws.cell(r, 6, total_cost); tc.number_format = NIS; tc.font = bold; tc.fill = gray; tc.alignment = center
for c in (7, 8):
    ws.cell(r, c).fill = gray
ts = ws.cell(r, 9, total_sell); ts.number_format = NIS; ts.font = bold; ts.fill = green; ts.alignment = center

# גיליון לוחות לקנייה
ps = wb.create_sheet("לוחות לקנייה")
ps.sheet_view.rightToLeft = True
ps.merge_cells("A1:C1")
h = ps.cell(1, 1, "לוחות לקנייה (תכנון נוחות)")
h.fill = navy; h.font = Font(bold=True, color="FFFFFF", size=14); h.alignment = center
for c, lab in enumerate(["פלטה (מ״מ)", "כמות לוחות", "עלות (₪)"], start=1):
    cell = ps.cell(2, c, lab); cell.fill = navy; cell.font = wbf
    cell.alignment = center; cell.border = border
    ps.column_dimensions[L(c)].width = 18
rr = 3
for s, c in sorted(conv_mix.items()):
    ps.cell(rr, 1, f"{s[0]}x{s[1]}").alignment = center
    ps.cell(rr, 2, c).alignment = center
    pc = ps.cell(rr, 3, c * PRICES[s]); pc.number_format = NIS; pc.alignment = center
    for cc in range(1, 4):
        ps.cell(rr, cc).border = border
    rr += 1
ps.cell(rr, 1, "סה״כ").font = bold; ps.cell(rr, 1).fill = gray; ps.cell(rr, 1).alignment = center
ps.cell(rr, 2, sum(conv_mix.values())).font = bold; ps.cell(rr, 2).fill = gray; ps.cell(rr, 2).alignment = center
tcc = ps.cell(rr, 3, conv_buy); tcc.number_format = NIS; tcc.font = bold; tcc.fill = gray; tcc.alignment = center
rr += 2
for k, (lab, val, fill) in enumerate([
        ("עלות חומר - תכנון נוחות:", conv_buy, None),
        ("עלות חומר - תכנון זול-ביותר:", cost_buy, None),
        ("תוספת עבור נוחות:", conv_buy - cost_buy, yellow),
        ("סה״כ הצעת מחיר (כולל 30%):", total_sell, green)]):
    lc = ps.cell(rr + k, 1, lab); lc.font = bold; lc.alignment = right
    ps.merge_cells(start_row=rr + k, start_column=1, end_row=rr + k, end_column=2)
    vc = ps.cell(rr + k, 3, val); vc.number_format = NIS; vc.font = bold; vc.alignment = center
    if fill:
        vc.fill = fill

OUT = "pricing.xlsx"
wb.save(OUT)
print(f"\nנשמר: {OUT}")
