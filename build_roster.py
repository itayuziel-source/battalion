#!/usr/bin/env python3
"""מחולל שבצק יציאות - קו הגנה סוריה.

מייצר קובץ אקסל (.xlsx) עם:
- צביעת שבת/ערב שבת (צהוב) וחגים (כתום) בכותרות התאריכים
- צביעה אוטומטית לתאים: יום בית (ירוק) / יום בסיס (תכלת) דרך עיצוב מותנה
- רשימה נפתחת לכל תא (בית / בסיס)
- שורות וסיכומים: סה"כ בבסיס / בבית לכל יום, וימי בית/בסיס לכל חייל
- גיליון מקרא נפרד

להארכת התקופה: שנה את START_DATE / END_DATE.
להוספת חגים: הוסף תאריכים ל-HOLIDAYS (פורמט date(שנה, חודש, יום)).
"""

from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- נתונים
START_DATE = date(2026, 6, 24)
END_DATE = date(2026, 9, 30)

# חגים (כתום) - ערב חג + החג עצמו. צומות וחוה"מ מושארים נייטרליים בכוונה.
HOLIDAYS: set[date] = {
    date(2026, 9, 11),  # ערב ראש השנה
    date(2026, 9, 12),  # ראש השנה א'
    date(2026, 9, 13),  # ראש השנה ב'
    date(2026, 9, 20),  # ערב יום כיפור
    date(2026, 9, 21),  # יום כיפור
    date(2026, 9, 25),  # ערב סוכות
    date(2026, 9, 26),  # סוכות א'
}

SOLDIERS = [
    ("איתי עוזיאל", "קלג"),
    ("אורי עוזיאל", "סמל נשקייה"),
    ("ינון איזון", "נשקייה"),
    ("שי ארנלדס", "נשקייה"),
    ("איתי שמש", "נשקייה"),
    ("אביעד כהן", "נשקייה"),
    ("אזוגי נתנאל", "אפסנאות"),
    ("ניב סורוקר", "אפסנאות"),
    ("צבי לונדין", "אפסנאות"),
    ("יוסף חגואל", "נשקייה"),
    ("אדם ממן", "סמל אפסנאות"),
    ("גבריאל יעקוביאן", "נשקייה"),
    ("יקיר אלימלך", "אפסנאות"),
    ("יוסי כהן", "נשקייה-רחפנים"),
    ("רון בן שושן", ""),
    ("דוד שמחי", ""),
    ("לב גינזבורג", "נשקייה"),
    ("יעקב גרמאי", "אפסנאות"),
    ("עידן אשטון", "אפסנאות"),
    ("עומר נוימן", ""),
    ("דוד עבאדה", "אפסנאות"),
]

# python weekday(): Mon=0..Sun=6  ->  אות עברית
HEB_DOW = {6: "א", 0: "ב", 1: "ג", 2: "ד", 3: "ה", 4: "ו", 5: "ש"}

# ---------------------------------------------------------------- צבעים
NAVY = "1F4E78"
HEADER_GRAY = "D9E1F2"
YELLOW = "FFFF00"      # שבת / ערב שבת
ORANGE = "FFC000"      # חג
HOME_GREEN = "C6EFCE"  # יום בית
BASE_BLUE = "BDD7EE"   # יום בסיס
TOTAL_FILL = "1F4E78"

fill_navy = PatternFill("solid", fgColor=NAVY)
fill_gray = PatternFill("solid", fgColor=HEADER_GRAY)
fill_yellow = PatternFill("solid", fgColor=YELLOW)
fill_orange = PatternFill("solid", fgColor=ORANGE)
fill_home = PatternFill("solid", fgColor=HOME_GREEN)
fill_base = PatternFill("solid", fgColor=BASE_BLUE)

white_bold = Font(bold=True, color="FFFFFF", size=11)
bold = Font(bold=True, size=11)
normal = Font(size=11)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")

thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def dates_range():
    d = START_DATE
    while d <= END_DATE:
        yield d
        d += timedelta(days=1)


DATES = list(dates_range())

# ---------------------------------------------------------------- בניית הגיליון
wb = Workbook()
ws = wb.active
ws.title = "שבצק יציאות"
ws.sheet_view.rightToLeft = True

# פריסה: A=שם B=תפקיד C=הערות  | D.. = תאריכים | אחרי כן: ימי בית / ימי בסיס
FIRST_DATE_COL = 4
n_dates = len(DATES)
home_sum_col = FIRST_DATE_COL + n_dates       # עמודת "ימי בית"
base_sum_col = home_sum_col + 1               # עמודת "ימי בסיס"
last_col = base_sum_col

TITLE_ROW = 1
MONTH_ROW = 2
HEAD_ROW = 3
FIRST_SOLDIER_ROW = 4

HEB_MONTH = {6: "יוני", 7: "יולי", 8: "אוגוסט", 9: "ספטמבר"}
last_soldier_row = FIRST_SOLDIER_ROW + len(SOLDIERS) - 1
base_total_row = last_soldier_row + 1
home_total_row = last_soldier_row + 2

# --- כותרת ראשית
ws.cell(TITLE_ROW, 1, "שבצק יציאות - קו הגנה סוריה")
ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=last_col)
c = ws.cell(TITLE_ROW, 1)
c.fill = fill_navy
c.font = Font(bold=True, color="FFFFFF", size=14)
c.alignment = center

# --- שורת חודשים (ממוזגת מעל עמודות התאריכים של כל חודש)
month_start_i = 0
for i in range(n_dates + 1):
    changed = i == n_dates or DATES[i].month != DATES[month_start_i].month
    if changed:
        c0 = FIRST_DATE_COL + month_start_i
        c1 = FIRST_DATE_COL + i - 1
        m = DATES[month_start_i]
        cell = ws.cell(MONTH_ROW, c0, f"{HEB_MONTH[m.month]} {m.year}")
        if c1 > c0:
            ws.merge_cells(start_row=MONTH_ROW, start_column=c0,
                           end_row=MONTH_ROW, end_column=c1)
        cell.fill = fill_navy
        cell.font = white_bold
        cell.alignment = center
        cell.border = border
        month_start_i = i

# --- כותרות עמודות קבועות
for col, label in ((1, "שם"), (2, "תפקיד"), (3, "הערות")):
    cell = ws.cell(HEAD_ROW, col, label)
    cell.fill = fill_navy
    cell.font = white_bold
    cell.alignment = center
    cell.border = border

# --- כותרות תאריכים (מספר + אות יום, צבע לפי שבת/חג)
for i, d in enumerate(DATES):
    col = FIRST_DATE_COL + i
    letter = HEB_DOW[d.weekday()]
    cell = ws.cell(HEAD_ROW, col, f"{d.day}\n{letter}")
    cell.alignment = center
    cell.border = border
    if d in HOLIDAYS:
        cell.fill = fill_orange
        cell.font = bold
    elif d.weekday() in (4, 5):  # שישי (ערב שבת) או שבת
        cell.fill = fill_yellow
        cell.font = bold
    else:
        cell.fill = fill_gray
        cell.font = bold
    ws.column_dimensions[get_column_letter(col)].width = 5

# --- כותרות סיכום לחייל
for col, label in ((home_sum_col, "ימי בית"), (base_sum_col, "ימי בסיס")):
    cell = ws.cell(HEAD_ROW, col, label)
    cell.fill = fill_navy
    cell.font = white_bold
    cell.alignment = center
    cell.border = border
    ws.column_dimensions[get_column_letter(col)].width = 9

first_date_letter = get_column_letter(FIRST_DATE_COL)
last_date_letter = get_column_letter(FIRST_DATE_COL + n_dates - 1)

# --- שורות חיילים
for r, (name, role) in enumerate(SOLDIERS, start=FIRST_SOLDIER_ROW):
    n = ws.cell(r, 1, name)
    n.font = bold
    n.alignment = right
    n.border = border
    rl = ws.cell(r, 2, role)
    rl.font = normal
    rl.alignment = center
    rl.border = border
    ws.cell(r, 3).border = border  # הערות
    for i in range(n_dates):
        ws.cell(r, FIRST_DATE_COL + i).border = border
    rng = f"{first_date_letter}{r}:{last_date_letter}{r}"
    h = ws.cell(r, home_sum_col)
    h.value = f'=COUNTIF({rng},"בית")'
    h.alignment = center
    h.border = border
    b = ws.cell(r, base_sum_col)
    b.value = f'=COUNTIF({rng},"בסיס")'
    b.alignment = center
    b.border = border

# --- שורות סה"כ לכל יום
for row, label, word in (
    (base_total_row, 'סה"כ בבסיס', "בסיס"),
    (home_total_row, 'סה"כ בבית', "בית"),
):
    lab = ws.cell(row, 1, label)
    lab.fill = fill_navy
    lab.font = white_bold
    lab.alignment = center
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    for i in range(n_dates):
        col = FIRST_DATE_COL + i
        letter = get_column_letter(col)
        cell = ws.cell(
            row, col,
            f'=COUNTIF({letter}{FIRST_SOLDIER_ROW}:{letter}{last_soldier_row},"{word}")',
        )
        cell.alignment = center
        cell.font = bold
        cell.border = border

# --- רשימה נפתחת + עיצוב מותנה על תאי חייל×תאריך
data_rng = (
    f"{first_date_letter}{FIRST_SOLDIER_ROW}:"
    f"{last_date_letter}{last_soldier_row}"
)
dv = DataValidation(type="list", formula1='"בית,בסיס"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(data_rng)
ws.conditional_formatting.add(
    data_rng, CellIsRule(operator="equal", formula=['"בית"'], fill=fill_home)
)
ws.conditional_formatting.add(
    data_rng, CellIsRule(operator="equal", formula=['"בסיס"'], fill=fill_base)
)

# --- מידות וקיבוע חלוניות
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 14
ws.row_dimensions[HEAD_ROW].height = 30
ws.freeze_panes = f"{first_date_letter}{FIRST_SOLDIER_ROW}"

# ---------------------------------------------------------------- גיליון מקרא
lg = wb.create_sheet("מקרא")
lg.sheet_view.rightToLeft = True
lg.column_dimensions["A"].width = 6
lg.column_dimensions["B"].width = 30
title = lg.cell(1, 2, "מקרא - שבצק יציאות")
title.font = Font(bold=True, size=13)
legend = [
    (fill_yellow, "שבת / ערב שבת"),
    (fill_orange, "חג"),
    (fill_home, "יום בית"),
    (fill_base, "יום בסיס"),
    (fill_gray, "יום רגיל (נייטרלי)"),
]
for i, (fill, label) in enumerate(legend, start=3):
    sw = lg.cell(i, 1)
    sw.fill = fill
    sw.border = border
    lg.cell(i, 2, label).font = normal

OUT = "shavtzak_yetziot.xlsx"
wb.save(OUT)
print(f"נשמר: {OUT} | חיילים: {len(SOLDIERS)} | ימים: {n_dates} "
      f"({START_DATE:%d.%m.%y}-{END_DATE:%d.%m.%y})")
