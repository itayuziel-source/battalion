#!/usr/bin/env python3
"""דוח ספירות ומלאי למפקד - נשקייה, 05/07/2026.

עיבוד של דוח CountsReportByUnitStock הגולמי לדוח מפקד:
- תמצית למפקד: כרטיסי מספרים, מבט לפי קטגוריה (עם פס אחוז חתימה),
  מימוש הקצאות לפי פלוגה, ודגשים אוטומטיים.
- פירוט פריטים: כל הפריטים מקובצים לקטגוריות, עם % חתום ופס נתונים.
- פילוח פלוגתי: הפריטים המוקצים לפלוגות - חתום/פנוי לכל פלוגה,
  הפנויים צבועים כדי שיקפצו לעין.
- נתוני מקור: הטבלה המקורית כפי שיוצאה מהמערכת.

סמנטיקת המקור: לכל פלוגה 'הקצאה' (כמה הוקצה) ו'פנוי' (כמה מההקצאה טרם
נחתם). 'כמות בקבלות' = חתום על חיילים; 'כמות מלאי' = לא חתום; סה"כ = שניהם.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.utils import get_column_letter as L

REPORT_DATE = "05/07/2026"
COMPANIES = ["פלוגה א", "פלוגה ב", "פלוגה ג", "מסייעת", "פלס\"ם"]

# (פריט, [פנוי,הקצאה x5 פלוגות], בקבלות, מלאי, סה"כ)
DATA = [
    ("M4A1", [18, 65, 24, 65, 25, 65, 4, 80, 7, 70], 267, 78, 345),
    ("כוונת איוטק", [16, 63, 21, 63, 23, 63, 6, 78, 0, 0], 262, 72, 334),
    ("M16", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 40, 132, 172),
    ("חד עיני - שחמ", [0, 16, 1, 15, 2, 20, 1, 15, 0, 0], 68, 5, 73),
    ("טריג", [3, 15, 0, 15, 0, 16, 0, 15, 0, 0], 66, 6, 72),
    ("קלע סער M4", [2, 15, 2, 15, 1, 16, 2, 15, 0, 0], 59, 11, 70),
    ("פק חום OGL", [8, 16, 2, 16, 2, 16, 8, 16, 0, 10], 56, 9, 65),
    ("פק שחור", [4, 10, 2, 10, 5, 10, 6, 10, 0, 8], 32, 26, 58),
    ("חד עיני - שחע", [9, 9, 2, 10, 4, 5, 2, 10, 0, 0], 21, 26, 47),
    ("דו עיני - עדי", [2, 7, 0, 7, 2, 7, 0, 6, 2, 10], 31, 6, 37),
    ("ליאור 4", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 34, 3, 37),
    ("M203 מטול", [0, 5, 0, 6, 0, 5, 0, 5, 0, 0], 25, 8, 33),
    ("M4 פלטופ קנה 1/12", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 23, 8, 31),
    ("כוונת השלכה מתקדמת", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 6, 23, 29),
    ("ציין לייזר לנגב", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 11, 14, 25),
    ("מאג", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 15, 10, 25),
    ("נגב דגם ב", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 15, 5, 20),
    ("מטל ממ", [2, 4, 3, 4, 0, 4, 4, 4, 0, 0], 7, 12, 19),
    ("דו עיני - עידו", [0, 2, 2, 2, 0, 2, 0, 2, 0, 8], 15, 1, 16),
    ("מכבים", [1, 3, 0, 3, 0, 3, 0, 3, 0, 1], 12, 1, 13),
    ("מכבית", [1, 3, 0, 3, 0, 3, 0, 3, 0, 0], 11, 1, 12),
    ("דו עיני - מיקרון", [0, 0, 0, 0, 0, 0, 0, 0, 0, 8], 12, 0, 12),
    ("רובצ M24", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 2, 6, 8),
    ("נגב 7.62", [1, 2, 0, 2, 0, 2, 0, 2, 0, 0], 7, 1, 8),
    ("משקפת תרמית - עמית", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 3, 3, 6),
    ("ליאור 4 פוקוס משתנה", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 5, 1, 6),
    ("נגב סער", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 3, 1, 4),
    ("ערמון", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 1, 1, 2),
    ("ליאור 3", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 2, 0, 2),
    ("מקלר", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 2, 0, 2),
    ("אורלי - מערכת גיל", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 2, 0, 2),
    ("משקפת תרמית - עמית איכון", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 1, 1, 2),
    ("נגב דגם ב  משופר", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0], 0, 1, 1),
]

CATEGORIES = [
    ("נשק אישי", ["M4A1", "M16", "קלע סער M4", "M4 פלטופ קנה 1/12", "M4 מקוצרר"]),
    ("מקלעים, מטולים וצלפים", ["נגב דגם ב", "נגב דגם ב  משופר", "נגב 7.62", "נגב סער",
                               "מאג", "מקלר", "M203 מטול", "מטל ממ", "רובצ M24"]),
    ("כוונות ולייזרים", ["כוונת איוטק", "טריג", "כוונת השלכה מתקדמת", "ליאור 3",
                          "ליאור 4", "ליאור 4 פוקוס משתנה", "ציין לייזר לנגב"]),
    ("אמצעי לילה ותרמי", ["חד עיני - שחמ", "חד עיני - שחע", "דו עיני - עדי",
                           "דו עיני - עידו", "דו עיני - מיקרון", "פק חום OGL",
                           "פק שחור", "מכבים", "מכבית", "משקפת תרמית - עמית",
                           "משקפת תרמית - עמית איכון"]),
    ("אחר", ["ערמון", "אורלי - מערכת גיל"]),
]

BY_NAME = {r[0]: r for r in DATA}
CAT_OF = {item: cat for cat, items in CATEGORIES for item in items}

# ---------------------------------------------------------------- סגנון
NAVY = "1F4E78"
BLUE2 = "2E75B6"
navy = PatternFill("solid", fgColor=NAVY)
sub = PatternFill("solid", fgColor=BLUE2)
gray = PatternFill("solid", fgColor="D9E1F2")
green = PatternFill("solid", fgColor="C6EFCE")
red = PatternFill("solid", fgColor="FFC7CE")
yellow = PatternFill("solid", fgColor="FFF2CC")
wbf = Font(bold=True, color="FFFFFF", size=11)
bold = Font(bold=True)
big = Font(bold=True, color="FFFFFF", size=16)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center", wrap_text=True)
thin = Side(style="thin", color="9AA7BD")
border = Border(thin, thin, thin, thin)
PCT = "0%"


def rtl(ws):
    ws.sheet_view.rightToLeft = True


def title(ws, text, span, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(1, 1, text)
    c.fill = navy; c.font = big; c.alignment = center
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
        s = ws.cell(2, 1, subtitle)
        s.font = Font(color="44546A", size=10); s.alignment = center


def head(ws, row, labels, widths=None, start=1, fill=None):
    for c, lab in enumerate(labels, start=start):
        cell = ws.cell(row, c, lab)
        cell.fill = fill or navy; cell.font = wbf; cell.alignment = center; cell.border = border
        if widths:
            ws.column_dimensions[L(c)].width = widths[c - start]


def databar(ws, rng):
    ws.conditional_formatting.add(rng, DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1,
        color="2E75B6", showValue=True))


# חישובים
def cat_sums():
    out = []
    for cat, items in CATEGORIES:
        rows = [BY_NAME[i] for i in items if i in BY_NAME]
        rec = sum(r[2] for r in rows); stk = sum(r[3] for r in rows); tot = sum(r[4] for r in rows)
        out.append((cat, tot, rec, stk))
    return out

def comp_sums():
    out = []
    for ci, comp in enumerate(COMPANIES):
        alloc = sum(r[1][ci * 2 + 1] for r in DATA)
        free = sum(r[1][ci * 2] for r in DATA)
        out.append((comp, alloc, alloc - free, free))
    return out

TOT = sum(r[4] for r in DATA)
REC = sum(r[2] for r in DATA)
STK = sum(r[3] for r in DATA)

wb = Workbook()

# ================================================================ תמצית למפקד
s = wb.active
s.title = "תמצית למפקד"
rtl(s)
SPAN = 10
title(s, f"דוח ספירות ומלאי – נשקייה | {REPORT_DATE}", SPAN,
      "מקור: דוח ספירות לפי פלוגה ומלאי | חתום = על קבלות אישיות | במלאי = טרם נחתם")
for c, w in zip(range(1, SPAN + 1), (21, 9, 9, 9, 11, 3, 13, 9, 9, 11)):
    s.column_dimensions[L(c)].width = w

# כרטיסי KPI
cards = [
    ("סה\"כ אמצעים", TOT, gray),
    ("חתום בקבלות", REC, green),
    ("במלאי (לא חתום)", STK, yellow),
    ("אחוז חתום", REC / TOT, gray),
]
s.row_dimensions[5].height = 30
for i, (lab, val, fill) in enumerate(cards):
    c0 = 1 + i * 2
    s.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 1)
    lc = s.cell(4, c0, lab); lc.fill = sub; lc.font = wbf; lc.alignment = center; lc.border = border
    s.cell(4, c0 + 1).border = border
    s.merge_cells(start_row=5, start_column=c0, end_row=5, end_column=c0 + 1)
    vc = s.cell(5, c0, val); vc.fill = fill; vc.font = Font(bold=True, size=18, color=NAVY)
    vc.alignment = center; vc.border = border
    s.cell(5, c0 + 1).border = border
    if lab == "אחוז חתום":
        vc.number_format = PCT

# מבט לפי קטגוריה
r0 = 7
s.cell(r0, 1, "מבט לפי קטגוריה").font = Font(bold=True, size=13)
head(s, r0 + 1, ["קטגוריה", "סה\"כ", "חתום", "במלאי", "% חתום"])
for i, (cat, tot, rec, stk) in enumerate(cat_sums(), start=1):
    r = r0 + 1 + i
    vals = [cat, tot, rec, stk, rec / tot if tot else 0]
    for c, v in enumerate(vals, start=1):
        cell = s.cell(r, c, v); cell.border = border
        cell.alignment = right if c == 1 else center
    s.cell(r, 5).number_format = PCT
last_cat = r0 + 1 + len(CATEGORIES)
tr = last_cat + 1
for c, v in enumerate(["סה\"כ", TOT, REC, STK, REC / TOT], start=1):
    cell = s.cell(tr, c, v); cell.font = bold; cell.border = border
    cell.alignment = right if c == 1 else center
    cell.fill = gray
s.cell(tr, 5).number_format = PCT
databar(s, f"E{r0 + 2}:E{tr}")

# מימוש הקצאות לפי פלוגה
r1 = tr + 2
s.cell(r1, 1, "מימוש הקצאות לפי פלוגה (פריטים מוקצים בלבד)").font = Font(bold=True, size=13)
head(s, r1 + 1, ["פלוגה", "הקצאה", "חתום", "פנוי", "% מימוש"])
for i, (comp, alloc, signed, free) in enumerate(comp_sums(), start=1):
    r = r1 + 1 + i
    vals = [comp, alloc, signed, free, signed / alloc if alloc else 0]
    for c, v in enumerate(vals, start=1):
        cell = s.cell(r, c, v); cell.border = border
        cell.alignment = right if c == 1 else center
    s.cell(r, 5).number_format = PCT
    if free:
        s.cell(r, 4).fill = yellow; s.cell(r, 4).font = bold
pr0, pr1 = r1 + 2, r1 + 1 + len(COMPANIES)
s.conditional_formatting.add(f"E{pr0}:E{pr1}",
    CellIsRule(operator="lessThan", formula=["0.75"], fill=red, font=bold))
s.conditional_formatting.add(f"E{pr0}:E{pr1}",
    CellIsRule(operator="between", formula=["0.75", "0.9"], fill=yellow))
s.conditional_formatting.add(f"E{pr0}:E{pr1}",
    CellIsRule(operator="greaterThanOrEqual", formula=["0.9"], fill=green))

# דגשים אוטומטיים
def make_highlights():
    hl = []
    # פלוגות עם רובים מוקצים ללא הקצאת איוטק
    m4 = BY_NAME["M4A1"][1]; io = BY_NAME["כוונת איוטק"][1]
    for ci, comp in enumerate(COMPANIES):
        if m4[ci * 2 + 1] > 0 and io[ci * 2 + 1] == 0:
            hl.append((f"ל{comp} מוקצים {m4[ci*2+1]} רובי M4A1 ללא הקצאת כוונות איוטק כלל.", True))
    # פריטים שרובם במלאי
    for name, comps, rec, stk, tot in DATA:
        if tot >= 20 and stk / tot >= 0.5:
            hl.append((f"{name}: {stk} מתוך {tot} ({stk*100//tot}%) יושבים במלאי ולא חתומים.", True))
    # הפלוגה עם המימוש הנמוך ביותר
    comps = [(c, a, sg, f) for c, a, sg, f in comp_sums() if a]
    worst = min(comps, key=lambda t: t[2] / t[1])
    hl.append((f"מימוש הקצאות נמוך ביותר: {worst[0]} – {worst[2]} חתומים מתוך {worst[1]} ({worst[2]*100//worst[1]}%), {worst[3]} פנויים.", False))
    total_free = sum(f for _, _, _, f in comp_sums())
    hl.append((f"סה\"כ {total_free} פריטים מוקצים לפלוגות וטרם נחתמו – מפורט בגיליון 'פילוח פלוגתי'.", False))
    return hl

r2 = pr1 + 2
s.cell(r2, 1, "דגשים").font = Font(bold=True, size=13)
for i, (txt, alert) in enumerate(make_highlights()):
    r = r2 + 1 + i
    s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=SPAN)
    c = s.cell(r, 1, ("⚠ " if alert else "• ") + txt)
    c.alignment = right
    if alert:
        c.font = Font(bold=True, color="9C0006")

# ================================================================ פירוט פריטים
d = wb.create_sheet("פירוט פריטים")
rtl(d)
heads = ["פריט", "חתום בקבלות", "במלאי", "סה\"כ", "% חתום", "הערה"]
title(d, "פירוט פריטים לפי קטגוריה", len(heads))
head(d, 3, heads, [26, 13, 10, 9, 12, 30])
r = 4
pct_rows = []
for cat, items in CATEGORIES:
    d.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(heads))
    hc = d.cell(r, 1, cat); hc.fill = sub; hc.font = wbf; hc.alignment = right
    r += 1
    rows = [BY_NAME[i] for i in items if i in BY_NAME]
    rows.sort(key=lambda t: -t[4])
    for name, comps, rec, stk, tot in rows:
        note = ""
        if tot and stk / tot >= 0.5 and tot >= 10:
            note = "רוב הכמות במלאי – לא חתום"
        elif stk == 0:
            note = "הכל חתום"
        vals = [name, rec, stk, tot, rec / tot if tot else 0, note]
        for c, v in enumerate(vals, start=1):
            cell = d.cell(r, c, v); cell.border = border
            cell.alignment = right if c in (1, 6) else center
        d.cell(r, 5).number_format = PCT
        if note.startswith("רוב"):
            d.cell(r, 6).fill = yellow; d.cell(r, 6).font = bold
        pct_rows.append(r)
        r += 1
    rec = sum(t[2] for t in rows); stk = sum(t[3] for t in rows); tot = sum(t[4] for t in rows)
    for c, v in enumerate([f"סה\"כ {cat}", rec, stk, tot, rec / tot if tot else 0, ""], start=1):
        cell = d.cell(r, c, v); cell.font = bold; cell.fill = gray; cell.border = border
        cell.alignment = right if c == 1 else center
    d.cell(r, 5).number_format = PCT
    r += 1
databar(d, f"E4:E{r - 1}")
d.freeze_panes = "A4"

# ================================================================ פילוח פלוגתי
p = wb.create_sheet("פילוח פלוגתי")
rtl(p)
alloc_rows = [row for row in DATA if any(row[1][i * 2 + 1] for i in range(5))]
ncols = 1 + 2 * len(COMPANIES) + 3
title(p, "פילוח פלוגתי – חתום / פנוי מתוך ההקצאה", ncols)
p.cell(2, 1, "פנוי = הוקצה לפלוגה וטרם נחתם על חייל. תאים צהובים/אדומים דורשים סגירה.").font = Font(color="44546A", size=10)
p.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
p.cell(2, 1).alignment = center
# כותרת כפולה
p.column_dimensions["A"].width = 22
p.cell(4, 1, "פריט").fill = navy; p.cell(4, 1).font = wbf; p.cell(4, 1).alignment = center
p.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
p.cell(3, 1).fill = navy; p.cell(3, 1).border = border; p.cell(4, 1).border = border
for ci, comp in enumerate(COMPANIES):
    c0 = 2 + ci * 2
    p.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 1)
    hc = p.cell(3, c0, comp); hc.fill = navy; hc.font = wbf; hc.alignment = center
    for off, lab in ((0, "חתום"), (1, "פנוי")):
        cell = p.cell(4, c0 + off, lab); cell.fill = sub; cell.font = wbf
        cell.alignment = center; cell.border = border
        p.column_dimensions[L(c0 + off)].width = 8
    p.cell(3, c0).border = border; p.cell(3, c0 + 1).border = border
c0 = 2 + len(COMPANIES) * 2
for off, lab, w in ((0, "סה\"כ הקצאה", 12), (1, "סה\"כ פנוי", 11), (2, "% מימוש", 10)):
    p.merge_cells(start_row=3, start_column=c0 + off, end_row=4, end_column=c0 + off)
    cell = p.cell(3, c0 + off, lab); cell.fill = navy; cell.font = wbf
    cell.alignment = center; cell.border = border
    p.cell(4, c0 + off).border = border
    p.column_dimensions[L(c0 + off)].width = w
r = 5
for name, comps, rec, stk, tot in sorted(alloc_rows, key=lambda t: -sum(t[1][i*2] for i in range(5))):
    p.cell(r, 1, name).alignment = right; p.cell(r, 1).border = border
    total_alloc = total_free = 0
    for ci in range(len(COMPANIES)):
        free, alloc = comps[ci * 2], comps[ci * 2 + 1]
        signed = alloc - free
        total_alloc += alloc; total_free += free
        sc = p.cell(r, 2 + ci * 2, signed if alloc else "")
        fc = p.cell(r, 3 + ci * 2, free if alloc else "")
        for cell in (sc, fc):
            cell.alignment = center; cell.border = border
        if alloc and free:
            fc.fill = red if free >= 10 else yellow
            fc.font = bold
    vals = [total_alloc, total_free, (total_alloc - total_free) / total_alloc if total_alloc else 0]
    for off, v in enumerate(vals):
        cell = p.cell(r, c0 + off, v); cell.alignment = center; cell.border = border
    p.cell(r, c0 + 2).number_format = PCT
    if total_free:
        p.cell(r, c0 + 1).fill = yellow; p.cell(r, c0 + 1).font = bold
    r += 1
p.conditional_formatting.add(f"{L(c0+2)}5:{L(c0+2)}{r-1}",
    CellIsRule(operator="lessThan", formula=["0.75"], fill=red, font=bold))
p.conditional_formatting.add(f"{L(c0+2)}5:{L(c0+2)}{r-1}",
    CellIsRule(operator="greaterThanOrEqual", formula=["0.9"], fill=green))
p.freeze_panes = "B5"

# ================================================================ נתוני מקור
raw = wb.create_sheet("נתוני מקור")
rtl(raw)
ncols = 4 + len(COMPANIES) * 2
title(raw, f"נתוני מקור – CountsReport {REPORT_DATE}", ncols)
raw.column_dimensions["A"].width = 24
raw.cell(3, 1, "שם פריט").fill = navy; raw.cell(3, 1).font = wbf; raw.cell(3, 1).alignment = center
raw.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
for ci, comp in enumerate(COMPANIES):
    c = 2 + ci * 2
    raw.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c + 1)
    hc = raw.cell(3, c, comp); hc.fill = navy; hc.font = wbf; hc.alignment = center
    for off, lab in ((0, "פנוי"), (1, "הקצאה")):
        cc = raw.cell(4, c + off, lab); cc.fill = sub; cc.font = wbf; cc.alignment = center; cc.border = border
        raw.column_dimensions[L(c + off)].width = 8
for off, lab in enumerate(["כמות בקבלות", "כמות מלאי", "סה\"כ"]):
    c = 2 + len(COMPANIES) * 2 + off
    raw.merge_cells(start_row=3, start_column=c, end_row=4, end_column=c)
    cc = raw.cell(3, c, lab); cc.fill = navy; cc.font = wbf; cc.alignment = center; cc.border = border
    raw.column_dimensions[L(c)].width = 12
for i, (name, comps, rec, stk, tot) in enumerate(DATA):
    r = 5 + i
    raw.cell(r, 1, name).alignment = right
    for j, v in enumerate(comps):
        raw.cell(r, 2 + j, v).alignment = center
    for off, v in enumerate([rec, stk, tot]):
        raw.cell(r, 2 + len(COMPANIES) * 2 + off, v).alignment = center
    for c in range(1, ncols + 1):
        raw.cell(r, c).border = border
raw.freeze_panes = "B5"

OUT = "count_report_2026-07-05.xlsx"
wb.save(OUT)
print("נשמר:", OUT, "| גיליונות:", wb.sheetnames)
