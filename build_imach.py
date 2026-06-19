# -*- coding: utf-8 -*-
"""
ממיר את אתר ה"ימ״ח — דשבורד מוכנות גדוד" לקובץ אקסל מסודר, עם דגש על
התיקים/הזיווד — מה נמצא איפה: זיווד רכבי ה-FMTV, זיווד 8 האמרים,
רשימות הנשק, ספר הקשר (דרכ"ש) ולוז המשימות.
"""
import re, json
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HTML="site_src/index.html"; OUT="ימח_דשבורד_גדוד.xlsx"
html=open(HTML,encoding="utf-8").read()
soup=BeautifulSoup(html,"lxml")

# ---------- צבעים/סגנון ----------
NAVY="1F3864"; BLUE="2E5496"; AMBER="B45309"; STEEL="8EAADB"; GREY="64748B"
WHITE="FFFFFF"; ZEBRA="F2F5FA"; CARD="0F172A"
thin=Side(style="thin",color="C7CED9")
BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(c): return PatternFill("solid",fgColor=c)
CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
CENV=Alignment(horizontal="center",vertical="center")
RGT=Alignment(horizontal="right",vertical="center",wrap_text=True)

wb=Workbook(); wb.remove(wb.active)

def new_sheet(title, headers, widths, header_fill=NAVY):
    ws=wb.create_sheet(title); ws.sheet_view.rightToLeft=True
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
    t=ws.cell(1,1,title); t.fill=fill(header_fill); t.font=Font(bold=True,color=WHITE,size=14); t.alignment=CENV
    ws.row_dimensions[1].height=28
    for i,h in enumerate(headers):
        c=ws.cell(2,i+1,h); c.fill=fill(header_fill); c.font=Font(bold=True,color=WHITE,size=11); c.alignment=CEN; c.border=BORDER
    ws.row_dimensions[2].height=24
    for i,w in enumerate(widths): ws.column_dimensions[get_column_letter(i+1)].width=w
    ws.freeze_panes="A3"; ws.auto_filter.ref=f"A2:{get_column_letter(len(headers))}2"
    return ws

def write_rows(ws, rows, start=3, bold_first=True):
    for k,row in enumerate(rows):
        r=start+k
        for i,v in enumerate(row):
            c=ws.cell(r,i+1, "" if v is None else v); c.border=BORDER
            c.alignment=RGT if (i==0 or isinstance(v,str) and len(str(v))>12) else CENV
            if bold_first and i==0: c.font=Font(bold=True)
            if k%2==1:
                if c.fill.fgColor.rgb in (None,"00000000"): c.fill=fill(ZEBRA)
    return start+len(rows)

# ============================================================
# 1) זיווד רכבים (FMTV) — window.DV
# ============================================================
raw=re.search(r'window\.DV=(\{.*\})\s*;?\s*$', [l for l in html.split("\n") if l.startswith("window.DV=")][0]).group(1)
DV=json.loads(re.sub(r'([{,])([a-z]):', r'\1"\2":', raw))
veh_rows=[]
for veh,info in DV.items():
    struct=info.get("t","")
    for comp in info.get("s",[]):
        pos=comp.get("p",""); typ=comp.get("d","")
        for item in comp.get("n",[]):
            veh_rows.append([veh, struct, pos, typ, item])
    if "x" in info:
        for item in info["x"].get("n",[]):
            veh_rows.append([veh, struct, info["x"].get("l","אחר"), "", item])
ws=new_sheet("זיווד רכבים (FMTV)",
    ["רכב","מבנה","עמדה / דולב","סוג עמדה","תוכן (חייל / תיק / פריט)"],
    [22,20,18,16,30])
write_rows(ws, veh_rows)

# ============================================================
# 2) זיווד האמרים — hammer-equip-container
# ============================================================
cont=soup.find(id="hammer-equip-container")
GROUP_MARK=("⬆","🔧","📡","📋","🎒","🧰","⚙","📦","🛠","🔩","🚪","🔦")
ham_rows=[]
for tog in cont.find_all(attrs={"onclick":True}):
    header=tog.get_text(" ",strip=True).replace("▼","").replace("▲","").strip()
    cz=re.search(r"צ׳?\s*([0-9]{5,})", header)
    cznum=cz.group(1) if cz else ""
    detail=tog.find_next_sibling()
    entries=[s.strip() for s in detail.stripped_strings if s.strip()] if detail else []
    group=""
    for e in entries:
        if any(e.startswith(m) or e[:2].strip().startswith(m) or m in e[:3] for m in GROUP_MARK):
            group=re.sub(r"^[^\wא-ת]+","",e).strip(); continue
        item, loc = (e.split(" - ",1)+[""])[:2] if " - " in e else (e,"")
        ham_rows.append([header, cznum, group, item.strip(), loc.strip()])
ws=new_sheet("זיווד האמרים",
    ["האמר","צ׳ נשק/רכב","קבוצת זיווד","פריט","מיקום / הערה"],
    [26,14,20,30,24], header_fill=AMBER)
write_rows(ws, ham_rows)

# ============================================================
# 3) נשקייה — רשימות נשק (tables 26-30)
# ============================================================
tables=soup.find_all("table")
wpn_rows=[]
for idx in range(26,31):
    rows=tables[idx].find_all("tr")
    for r in rows[1:]:
        cells=[c.get_text(" ",strip=True) for c in r.find_all(["td","th"])]
        if not any(cells): continue
        ident=cells[0] if cells else ""
        grp=re.match(r"[A-Za-z]+", ident)
        listname=f"רשימה {grp.group(0)}" if grp else "רשימה"
        wpn_rows.append([listname]+ (cells+[""]*6)[:6])
ws=new_sheet("נשקייה (רשימות נשק)",
    ["רשימה","#","סוג","צ׳ נשק","כוונת","שם","סמן"],
    [12,10,14,14,14,20,14], header_fill=BLUE)
write_rows(ws, wpn_rows, bold_first=False)

# ============================================================
# 4) אנשי קשר — דרכ"ש (tables 0-24)
# ============================================================
con_rows=[]
for idx in range(0,25):
    rows=tables[idx].find_all("tr")
    for r in rows[1:]:
        cells=[c.get_text(" ",strip=True) for c in r.find_all(["td","th"])]
        if not any(cells): continue
        con_rows.append([idx+1]+(cells+[""]*7)[:7])
ws=new_sheet("אנשי קשר (דרכ\"ש)",
    ["בלוק","תפקיד","מסגרת","שם","נייד","סל. צ.","אדום","מטכלי"],
    [8,22,18,20,16,14,14,14], header_fill=GREY)
write_rows(ws, con_rows, bold_first=False)

# ============================================================
# 5) לוז משימות — table 25
# ============================================================
t25=tables[25]; rows25=t25.find_all("tr")
hdr25=[c.get_text(" ",strip=True) for c in rows25[0].find_all(["td","th"])]
task_rows=[]
for r in rows25[1:]:
    cells=[c.get_text(" ",strip=True) for c in r.find_all(["td","th"])]
    if any(cells): task_rows.append((cells+[""]*len(hdr25))[:len(hdr25)])
ws=new_sheet("לוז משימות", hdr25, [40]+[10]*(len(hdr25)-1), header_fill=BLUE)
write_rows(ws, task_rows)

# ============================================================
# 0) סקירה — דף שער
# ============================================================
ov=wb.create_sheet("סקירה",0); ov.sheet_view.rightToLeft=True; ov.sheet_view.showGridLines=False
ov.merge_cells("A1:D1")
h=ov.cell(1,1,"ימ״ח — דשבורד מוכנות גדוד  |  גיבוי נתונים לאקסל")
h.font=Font(bold=True,size=16,color=WHITE); h.fill=fill(NAVY); h.alignment=CENV; ov.row_dimensions[1].height=34
ov.merge_cells("A2:D2")
ov.cell(2,1,"הופק מתוך האתר apricot-gusty-43.tiiny.site — כל הנתונים נשמרו בלשוניות נפרדות").font=Font(italic=True,color="595959")
cards=[("זיווד רכבים (FMTV)",f"{len(DV)} רכבים · {len(veh_rows)} פריטים"),
       ("זיווד האמרים",f"8 האמרים · {len(ham_rows)} פריטים"),
       ("רשימות נשק",f"{len(wpn_rows)} כלים"),
       ("אנשי קשר (דרכ\"ש)",f"{len(con_rows)} אנשי קשר"),
       ("לוז משימות",f"{len(task_rows)} משימות")]
r=4
ov.cell(r,1,"לשונית").font=Font(bold=True,color=WHITE); ov.cell(r,1).fill=fill(BLUE)
ov.cell(r,2,"תוכן").font=Font(bold=True,color=WHITE); ov.cell(r,2).fill=fill(BLUE)
for c in (1,2): ov.cell(r,c).alignment=CEN; ov.cell(r,c).border=BORDER
for i,(name,desc) in enumerate(cards):
    rr=r+1+i
    a=ov.cell(rr,1,name); a.font=Font(bold=True); a.alignment=RGT; a.border=BORDER
    b=ov.cell(rr,2,desc); b.alignment=RGT; b.border=BORDER
ov.column_dimensions["A"].width=28; ov.column_dimensions["B"].width=34
ov.merge_cells(f"A{r+1+len(cards)+1}:D{r+1+len(cards)+1}")
note=ov.cell(r+1+len(cards)+1,1,"הערה: \"זיווד רכבים\" ו\"זיווד האמרים\" הם ה\"מה נמצא איפה\" — כל פריט/תיק והעמדה בה הוא מאוחסן.")
note.alignment=Alignment(horizontal="right",wrap_text=True); note.font=Font(italic=True,color="595959")

wb.active=0
wb.save(OUT)
print("נשמר:",OUT)
print(f"FMTV פריטים={len(veh_rows)} | האמרים פריטים={len(ham_rows)} | נשק={len(wpn_rows)} | קשר={len(con_rows)} | משימות={len(task_rows)}")
