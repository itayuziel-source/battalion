# -*- coding: utf-8 -*-
"""
מזריק ערכים מחושבים (<v>) לכל תא-נוסחה בקובץ, כך שהמספרים יוצגו בכל מציג
(אקסל, גוגל, תצוגה מקדימה) ולא רק אחרי חישוב-מחדש — בלי לפגוע בנוסחאות.
"""
import zipfile, re, shutil
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import get_column_letter

ORIG="sefira_input.xlsx"; FILE="שבצק_יציאות_מסודר.xlsx"; ROSTER="שבצק יציאות"; TS='ת״ש'
MAIN="{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_R="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

o=openpyxl.load_workbook(ORIG,data_only=True)
f=openpyxl.load_workbook(ORIG,data_only=False)

# 1) ערכים שמורים מקוריים לכל תא-נוסחה (לכל לשונית)
cache={}
for name in f.sheetnames:
    wf=f[name]; wo=o[name]; d={}
    for row in wf.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith("="):
                v=wo[c.coordinate].value
                if v is not None: d[c.coordinate]=v
    cache[name]=d

# 2) חישוב ערכי בלוק הספירה (לשונית השבצק)
wsd=o[ROSTER]
def st(r,c):
    v=wsd.cell(r,c).value
    return "" if v is None else str(v).strip()
roles={r:str(wsd.cell(r,2).value or "") for r in range(5,31)}
n_names=sum(1 for r in range(5,31) if str(wsd.cell(r,1).value or "").strip())
reasons={44:"חופשה",45:TS,46:"קורס",47:"גימלים",48:"מבחן",49:"חול",50:"בית",51:"X",52:"?"}
comp={}
for c in range(4,106):
    Lc=get_column_letter(c); col_st=[st(r,c) for r in range(5,31)]
    base=sum(1 for v in col_st if v=="בסיס"); nonempty=sum(1 for v in col_st if v!="")
    comp[f"{Lc}31"]=base
    comp[f"{Lc}32"]=sum(1 for r in range(5,31) if "מחתים" in roles[r] and st(r,c)=="בסיס")
    comp[f"{Lc}33"]=sum(1 for r in range(5,31) if "לוגיסטיקה" in roles[r] and st(r,c)=="בסיס")
    comp[f"{Lc}34"]=sum(1 for v in col_st if v=="בית")
    comp[f"{Lc}35"]=nonempty-base
    for rr,key in reasons.items(): comp[f"{Lc}{rr}"]=sum(1 for v in col_st if v==key)
    comp[f"{Lc}53"]=n_names-nonempty

roster_vals=dict(cache[ROSTER]); roster_vals.update(comp)
vals_by_sheet={ROSTER:roster_vals}
for name in f.sheetnames:
    if name!=ROSTER: vals_by_sheet[name]=cache[name]

# 3) מיפוי שם-לשונית -> קובץ worksheet xml
z=zipfile.ZipFile(FILE)
root=ET.fromstring(z.read("xl/workbook.xml").decode("utf-8"))
name_to_rid={sh.get("name"):sh.get(NS_R+"id") for sh in root.iter(MAIN+"sheet")}
rels=ET.fromstring(z.read("xl/_rels/workbook.xml.rels").decode("utf-8"))
rid_to_t={rel.get("Id"):rel.get("Target") for rel in rels}
def path_for(name):
    t=rid_to_t[name_to_rid[name]].lstrip("/")
    return t if t.startswith("xl/") else "xl/"+t
z.close()

def inject(xml, vals):
    cnt=0
    for ref,val in vals.items():
        if isinstance(val,float) and val.is_integer(): val=int(val)
        pat=re.compile(r'(<c r="'+re.escape(ref)+r'"[^>]*>)(<f[^>]*>.*?</f>)(<v\s*/>|<v></v>)(</c>)')
        xml,n=pat.subn(lambda m: m.group(1)+m.group(2)+f"<v>{val}</v>"+m.group(4), xml, count=1)
        cnt+=n
    return xml,cnt

zin=zipfile.ZipFile(FILE,"r")
modified={}; total=0
for name,vals in vals_by_sheet.items():
    p=path_for(name)
    newxml,cnt=inject(zin.read(p).decode("utf-8"), vals)
    modified[p]=newxml; total+=cnt
    print(f"{name}: injected {cnt} values")
tmp=FILE+".tmp"
zout=zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data=zin.read(item.filename)
    if item.filename in modified: data=modified[item.filename].encode("utf-8")
    zout.writestr(item,data)
zin.close(); zout.close(); shutil.move(tmp,FILE)
print("total injected:",total)
