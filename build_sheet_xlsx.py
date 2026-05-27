#!/usr/bin/env python3
"""מפיק אקסל עם חישוב כמות לוחות הפח (ללא חפיפות)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from sheet_calc import PANELS, STOCKS, split_panel, pack

REQ_AREA = sum(w * h * q for w, h, q in PANELS)  # מ"מ^2

# חישוב סיכומים
totals = []
for sw, sh in STOCKS:
    pieces = []
    for w, h, q in PANELS:
        pieces.extend(split_panel(w, h, sw, sh) * q)
    n = pack(pieces, sw, sh)
    area = n * sw * sh
    totals.append((sw, sh, n, area, 1 - REQ_AREA / area))
best = min(totals, key=lambda r: r[4])

# סגנון
NAVY = "1F4E78"
fill_navy = PatternFill("solid", fgColor=NAVY)
fill_green = PatternFill("solid", fgColor="C6EFCE")
fill_gray = PatternFill("solid", fgColor="D9E1F2")
white_bold = Font(bold=True, color="FFFFFF", size=11)
bold = Font(bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")
thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, labels, widths):
    for c, (lab, w) in enumerate(zip(labels, widths), start=1):
        cell = ws.cell(row, c, lab)
        cell.fill = fill_navy; cell.font = white_bold
        cell.alignment = center; cell.border = border
        ws.column_dimensions[L(c)].width = w


def title(ws, text, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    t = ws.cell(1, 1, text)
    t.fill = fill_navy; t.font = Font(bold=True, color="FFFFFF", size=14); t.alignment = center


wb = Workbook()

# ---------------- סיכום
s = wb.active
s.title = "סיכום"
s.sheet_view.rightToLeft = True
title(s, "חישוב כמות לוחות פח - ללא חפיפות", 4)
header(s, 2, ["מידת לוח (מ\"מ)", "לוחות נדרשים", "שטח כולל (מ\"ר)", "פחת"], [16, 14, 16, 10])
for i, (sw, sh, n, area, waste) in enumerate(totals, start=3):
    s.cell(i, 1, f"{sw}x{sh}").alignment = center
    s.cell(i, 2, n).alignment = center
    ca = s.cell(i, 3, round(area / 1e6, 2)); ca.alignment = center
    cw = s.cell(i, 4, round(waste, 4)); cw.alignment = center; cw.number_format = "0%"
    fill = fill_green if (sw, sh) == (best[0], best[1]) else None
    for c in range(1, 5):
        cell = s.cell(i, c); cell.border = border
        if fill:
            cell.fill = fill; cell.font = bold
r = 3 + len(totals) + 1
s.cell(r, 1, "שטח נדרש נטו (מ\"ר):").alignment = right; s.cell(r, 1).font = bold
s.cell(r, 2, round(REQ_AREA / 1e6, 2)).alignment = center
s.cell(r + 2, 1, "המלצה:").font = bold
s.cell(r + 2, 2, f"{best[0]}x{best[1]} → {best[2]} לוחות, פחת {best[4]*100:.0f}%").font = bold
notes = [
    "הנחות:",
    "• תפרים מותרים (פאנל מורכב מכמה לוחות).",
    "• ללא חפיפות וללא קרף חיתוך.",
    "• ניצול שאריות בין פאנלים.",
    "• סיבוב חתיכות חופשי (אם לפח יש כיוון - יש לחשב מחדש).",
]
for k, txt in enumerate(notes):
    s.cell(r + 4 + k, 1, txt).alignment = right

# ---------------- רשימת פאנלים
p = wb.create_sheet("רשימת פאנלים")
p.sheet_view.rightToLeft = True
title(p, "רשימת פאנלים", 6)
header(p, 2, ["#", "מידות (מ\"מ)", "אורך", "רוחב", "כמות", "שטח כולל (מ\"ר)"],
       [5, 16, 10, 10, 8, 16])
tot_area = 0
for i, (w, h, q) in enumerate(PANELS, start=1):
    row = 2 + i
    a = w * h * q / 1e6
    tot_area += a
    vals = [i, f"{w}x{h}", w, h, q, round(a, 2)]
    for c, v in enumerate(vals, start=1):
        cell = p.cell(row, c, v); cell.alignment = center; cell.border = border
trow = 3 + len(PANELS)
p.cell(trow, 1, "סה\"כ").font = bold
p.merge_cells(start_row=trow, start_column=1, end_row=trow, end_column=5)
p.cell(trow, 1).alignment = center; p.cell(trow, 1).fill = fill_gray
tc = p.cell(trow, 6, round(tot_area, 2)); tc.font = bold; tc.alignment = center; tc.fill = fill_gray

# ---------------- פירוט לפי סוג (לוח מומלץ)
bw, bh = best[0], best[1]
d = wb.create_sheet(f"פירוט {bw}x{bh}")
d.sheet_view.rightToLeft = True
title(d, f"פירוט לפי סוג פאנל - לוח {bw}x{bh}", 4)
header(d, 2, ["פאנל (מ\"מ)", "כמות", "חתיכות לפאנל", "לוחות (עצמאי)"], [16, 8, 14, 14])
sum_indep = 0
for i, (w, h, q) in enumerate(PANELS, start=1):
    row = 2 + i
    n_pieces = len(split_panel(w, h, bw, bh))
    indep = pack(split_panel(w, h, bw, bh) * q, bw, bh)
    sum_indep += indep
    for c, v in enumerate([f"{w}x{h}", q, n_pieces, indep], start=1):
        cell = d.cell(row, c, v); cell.alignment = center; cell.border = border
nrow = 3 + len(PANELS) + 1
d.cell(nrow, 1, f"סכום עצמאי: {sum_indep} לוחות").font = bold
d.merge_cells(start_row=nrow, start_column=1, end_row=nrow, end_column=4)
d.cell(nrow + 1, 1,
       f"בפועל עם ניצול שאריות משותף: {best[2]} לוחות (פחות מהסכום העצמאי).").font = bold
d.merge_cells(start_row=nrow + 1, start_column=1, end_row=nrow + 1, end_column=4)

OUT = "sheet_metal_calc.xlsx"
wb.save(OUT)
print(f"נשמר: {OUT} | המלצה: {bw}x{bh} -> {best[2]} לוחות, פחת {best[4]*100:.0f}%")
