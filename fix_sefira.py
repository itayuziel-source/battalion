# -*- coding: utf-8 -*-
"""
מתקן ומסדר את "הספירה למטה" בשבצק היציאות:
- ממלא את נוסחאות הספירה לאורך כל 102 ימי החודשים (D:DA) — היום היו רק בעמודה הראשונה
- מתקן את ספירת התפקידים לתפקידים שקיימים בפועל (מחתים / לוגיסטיקה)
- מוסיף "סה"כ לא בבסיס" ופילוח 'לא בבסיס' לפי סיבה לכל יום
- העיצוב המותנה (אדום מתחת לסף מצבה) כבר קיים בקובץ ויידלק אוטומטית
"""
import openpyxl
from copy import copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

IN="sefira_input.xlsx"; OUT="שבצק_יציאות_מסודר.xlsx"
SH="שבצק יציאות"
FIRST_DAY=4; LAST_DAY=105            # D..DA
S_FIRST, S_LAST = 5, 30              # שורות החיילים (כולל שורת חיץ 30)
TS='ת״ש'                              # עם גרשיים עבריים (U+05F4) כמו במקור

wb=openpyxl.load_workbook(IN)
wb.calculation=CalcProperties(fullCalcOnLoad=True)   # חישוב מחדש בפתיחה
ws=wb[SH]

def col(c): return get_column_letter(c)
def setf(r,c,formula,style_src=None):
    cell=ws.cell(r,c,formula)
    if style_src is not None: cell._style=copy(style_src._style)
    return cell

# סגנון מקור להעתקה (התא הקיים D31)
src31=ws.cell(31,FIRST_DAY)

# ---------- 1) מילוי נוסחאות הספירה לכל יום ----------
for c in range(FIRST_DAY, LAST_DAY+1):
    Lc=col(c); rng=f"{Lc}{S_FIRST}:{Lc}{S_LAST}"
    setf(31,c, f'=COUNTIF({rng},"בסיס")', src31)                                  # סה"כ בבסיס
    setf(32,c, f'=COUNTIFS($B${S_FIRST}:$B${S_LAST},"*מחתים*",{rng},"בסיס")', src31)   # מחתים בבסיס
    setf(33,c, f'=COUNTIFS($B${S_FIRST}:$B${S_LAST},"*לוגיסטיקה*",{rng},"בסיס")', src31)# לוגיסטיקה בבסיס
    setf(34,c, f'=COUNTIF({rng},"בית")', src31)                                   # סה"כ בבית
    setf(35,c, f'=COUNTA({rng})-COUNTIF({rng},"בסיס")', src31)                    # סה"כ לא בבסיס

# ---------- 2) תוויות מעודכנות ----------
ws.cell(32,1,"מחתים בבסיס")
ws.cell(33,1,"לוגיסטיקה בבסיס")
ws.cell(40,1,"מינימום מחתים")
ws.cell(41,1,"מינימום לוגיסטיקה")
# שורה 35 — תווית חדשה "סה"כ לא בבסיס" (ממוזג A:C כמו השורות מעל)
try: ws.merge_cells("A35:C35")
except Exception: pass
lbl35=ws.cell(35,1,"סה\"כ לא בבסיס")
lbl34=ws.cell(34,1)
lbl35._style=copy(lbl34._style)

# ---------- 3) פילוח 'לא בבסיס' לפי סיבה לכל יום ----------
thin=Side(style="thin",color="BFBFBF")
BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
HDRF=PatternFill("solid",fgColor="1F3864")
LBLF=PatternFill("solid",fgColor="D9E1F2")
CEN=Alignment(horizontal="center",vertical="center")
RGT=Alignment(horizontal="right",vertical="center")

HR=43  # שורת כותרת הפילוח
ws.merge_cells(start_row=HR,start_column=1,end_row=HR,end_column=3)
hc=ws.cell(HR,1,"פילוח 'לא בבסיס' לפי סיבה — לכל יום")
hc.fill=HDRF; hc.font=Font(bold=True,color="FFFFFF",size=11); hc.alignment=RGT
# ערך מספרי לכותרת לאורך הימים? לא — רק תאי הסיבות
reasons=[
 ("חופשה",  lambda rng:f'=COUNTIF({rng},"חופשה")'),
 (TS,       lambda rng:f'=COUNTIF({rng},"{TS}")'),
 ("קורס",   lambda rng:f'=COUNTIF({rng},"קורס")'),
 ("גימלים", lambda rng:f'=COUNTIF({rng},"גימלים")'),
 ("מבחן",   lambda rng:f'=COUNTIF({rng},"מבחן")'),
 ("חול",    lambda rng:f'=COUNTIF({rng},"חול")'),
 ("בית",    lambda rng:f'=COUNTIF({rng},"בית")'),
 ("X",      lambda rng:f'=COUNTIF({rng},"X")'),
 ("? (לא ידוע)", lambda rng:f'=COUNTIF({rng},"~?")'),   # ~? = סימן שאלה מילולי
 ("ללא סטטוס (ריק)", lambda rng:f'=COUNTA($A${S_FIRST}:$A${S_LAST})-COUNTA({rng})'),
]
for i,(label,fn) in enumerate(reasons):
    r=HR+1+i
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
    lc=ws.cell(r,1,label); lc.fill=LBLF; lc.font=Font(bold=True); lc.alignment=RGT; lc.border=BORDER
    for c in range(FIRST_DAY,LAST_DAY+1):
        Lc=col(c); rng=f"{Lc}{S_FIRST}:{Lc}{S_LAST}"
        cell=ws.cell(r,c,fn(rng)); cell.alignment=CEN; cell.border=BORDER

wb.save(OUT)
print("נשמר:",OUT)
print(f"מולאו ימי ספירה: עמודות {col(FIRST_DAY)}..{col(LAST_DAY)} | שורות 31-35 + פילוח {HR+1}-{HR+len(reasons)}")
