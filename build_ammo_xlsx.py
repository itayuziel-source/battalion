#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the shared battalion ammunition-management workbook (ammo_management.xlsx).

Everyone edits one sheet only — 'יומן תנועות' — and the dashboard recomputes:
total signed = container (Hatzerot Yosef) + with forces in Lebanon + consumed.
Opening balances, companies and locations are seeded from the source file the
user uploaded (בילו + חצרות יסף signatures, July 2026).
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LOG_MAX = 3000          # last log row covered by every formula / dropdown
OLIVE = "55663A"
OLIVE_SOFT = "E6E9DC"
PAPER = "F5F4EF"
INK2 = "5B5E52"
GREEN = "1BAF7A"
BLUE = "2A78D6"
ORANGE = "EB6834"
RED_TXT = "B83A3A"
RED_FILL = "F6E4E4"

F = lambda **kw: Font(name="Arial", **kw)
THIN = Side(style="thin", color="D8D6CB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RIGHT = Alignment(horizontal="right", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")

# ---------------------------------------------------------------- source data
CATS = ["פצצות מרגמה 120", 'נ"ט', "מטולים", "רימונים", "תחמושת קלה", "אחר"]
ITEMS = [  # (name, category)
    ('פצמ"ר 120 נפיץ', CATS[0]), ('פצמ"ר 120 עשן', CATS[0]), ('פצמ"ר 120 תאורה', CATS[0]),
    ("מטאדור", CATS[1]), ("לאו", CATS[1]), ("גיל יום", CATS[1]), ("גיל סימן 2", CATS[1]),
    ("מטול נפיץ", CATS[2]), ("מטול תאורה", CATS[2]), ("מטול עשן", CATS[2]), ('מקל"ר', CATS[2]),
    ("רימון רסס 26ב'", CATS[3]), ("רימון עשן אפור", CATS[3]), ("רימון עשן כחול", CATS[3]),
    ("רימון עשן צהוב", CATS[3]), ("רימון עשן אדום", CATS[3]), ("רימון עשן ירוק", CATS[3]),
    ("7.62 משורשר", CATS[4]), ("5.56 משורשר", CATS[4]), ("5.56 דור ב'", CATS[4]), ("ליאה", CATS[4]),
]
COMPANIES = [
    ("א", ["חניתה"]),
    ("ב", ["מגדל זון", "טיר חרפא"]),
    ("ג", ["מגדל זון"]),
    ("מסייעת", ["מרגמות", "אוהד"]),
    ('חפ"ק מג"ד', ["מגדל זון"]),
    ("פתן", ["מגדל זון"]),
    ("שריון", ["טיר חרפא", "מגדל זון"]),
]
T_INTAKE, T_ISSUE, T_CONSUME, T_RETURN = "חתימת מלאי", "ניפוק ללבנון", "דיווח ירי", "החזרה למכולה"
TYPES = [T_INTAKE, T_ISSUE, T_CONSUME, T_RETURN]

# opening intakes from the uploaded source file: (item, qty, source, date, signed by)
OPENING = [
    ('פצמ"ר 120 נפיץ', 192, "בילו", None, ""), ('פצמ"ר 120 עשן', 37, "בילו", None, ""),
    ('פצמ"ר 120 תאורה', 36, "בילו", None, ""), ("מטאדור", 4, "בילו", None, ""),
    ("לאו", 18, "בילו", None, ""), ("גיל יום", 5, "בילו", None, ""), ("גיל סימן 2", 2, "בילו", None, ""),
    ("מטול נפיץ", 481, "בילו", None, ""), ("מטול תאורה", 575, "בילו", None, ""),
    ("מטול עשן", 357, "בילו", None, ""), ("רימון רסס 26ב'", 624, "בילו", None, ""),
    ("7.62 משורשר", 70100, "בילו", None, ""), ("5.56 משורשר", 72200, "בילו", None, ""),
    ("5.56 דור ב'", 299200, "בילו", None, ""),
    ("רימון עשן אפור", 88, "בילו", None, ""), ("רימון עשן כחול", 80, "בילו", None, ""),
    ("רימון עשן צהוב", 75, "בילו", None, ""), ("רימון עשן אדום", 80, "בילו", None, ""),
    ("רימון עשן ירוק", 80, "בילו", None, ""),
    ("ליאה", 12000, "חצרות יסף", None, ""),
    ('מקל"ר', 192, "חצרות יסף", date(2026, 7, 19), "אדם ממן"),
    ("7.62 משורשר", 22160, "חצרות יסף", date(2026, 7, 19), "אורי עוזיאל"),
]

wb = Workbook()

# ================================================================ הגדרות
st = wb.active
st.title = "הגדרות"
st.sheet_view.rightToLeft = True
st["A1"], st["B1"], st["C1"], st["D1"], st["E1"] = "פריט", "קטגוריה", "פלוגה", "מיקום", "סוג תנועה"
for c in "ABCDE":
    st[f"{c}1"].font = F(bold=True, size=11)
    st[f"{c}1"].fill = PatternFill("solid", fgColor=OLIVE_SOFT)
    st[f"{c}1"].alignment = RIGHT
for i, (name, cat) in enumerate(ITEMS, start=2):
    st[f"A{i}"], st[f"B{i}"] = name, cat
row = 2
ALL_LOCS = []
for co, locs in COMPANIES:
    st[f"C{row}"] = co
    row += 1
    for l in locs:
        if l not in ALL_LOCS:
            ALL_LOCS.append(l)
for i, l in enumerate(ALL_LOCS, start=2):
    st[f"D{i}"] = l
for i, t in enumerate(TYPES, start=2):
    st[f"E{i}"] = t
for col, w in zip("ABCDE", (20, 18, 14, 14, 16)):
    st.column_dimensions[col].width = w
for r in st.iter_rows(min_row=2, max_row=40, max_col=5):
    for c in r:
        c.font = F(size=11)
        c.alignment = RIGHT
st["G1"] = "שורות פנויות בכל רשימה מיועדות להוספת פריטים/פלוגות/מיקומים חדשים — הרשימות הנפתחות ביומן קוראות עד שורה 40."
st["G1"].font = F(size=10, italic=True, color=INK2)

# ================================================================ יומן תנועות
lg = wb.create_sheet("יומן תנועות")
lg.sheet_view.rightToLeft = True
headers = ["סוג תנועה", "תאריך", "פריט", "כמות", "פלוגה", "מיקום", "מקור", "שם מלא", "הערות"]
for j, h in enumerate(headers, start=1):
    c = lg.cell(row=1, column=j, value=h)
    c.font = F(bold=True, size=11, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=OLIVE)
    c.alignment = CENTER
    c.border = BOX
for j, w in enumerate((16, 12, 20, 11, 12, 13, 12, 18, 28), start=1):
    lg.column_dimensions[get_column_letter(j)].width = w
lg.freeze_panes = "A2"

for i, (item, qty, src, d, person) in enumerate(OPENING, start=2):
    lg.cell(row=i, column=1, value=T_INTAKE)
    if d:
        dc = lg.cell(row=i, column=2, value=d)
        dc.number_format = "DD.MM.YYYY"
    lg.cell(row=i, column=3, value=item)
    lg.cell(row=i, column=4, value=qty).number_format = "#,##0"
    lg.cell(row=i, column=7, value=src)
    lg.cell(row=i, column=8, value=person)
    lg.cell(row=i, column=9, value="יתרת פתיחה מקובץ המקור")
for r in lg.iter_rows(min_row=2, max_row=LOG_MAX, max_col=9):
    for c in r:
        c.font = F(size=11)
        c.alignment = RIGHT
        if c.column == 4:
            c.number_format = "#,##0"
        if c.column == 2:
            c.number_format = "DD.MM.YYYY"

def dv(formula, rng, msg):
    d = DataValidation(type="list", formula1=formula, allow_blank=True,
                       showErrorMessage=True, errorTitle="ערך לא ברשימה", error=msg)
    lg.add_data_validation(d)
    d.add(rng)

dv("=הגדרות!$E$2:$E$5", f"A2:A{LOG_MAX}", "בחרו סוג תנועה מהרשימה")
dv("=הגדרות!$A$2:$A$40", f"C2:C{LOG_MAX}", "בחרו פריט מהרשימה (או הוסיפו אותו קודם בגיליון הגדרות)")
dv("=הגדרות!$C$2:$C$40", f"E2:E{LOG_MAX}", "בחרו פלוגה מהרשימה (או הוסיפו אותה קודם בגיליון הגדרות)")
dv("=הגדרות!$D$2:$D$40", f"F2:F{LOG_MAX}", "בחרו מיקום מהרשימה (או הוסיפו אותו קודם בגיליון הגדרות)")

# highlight incomplete rows: type filled but qty/item missing
lg.conditional_formatting.add(
    f"A2:I{LOG_MAX}",
    FormulaRule(formula=['AND($A2<>"",OR($C2="",$D2=""))'],
                fill=PatternFill("solid", fgColor=RED_FILL), stopIfTrue=False),
)

# ================================================================ דשבורד
db = wb.create_sheet("דשבורד", 0)
db.sheet_view.rightToLeft = True
db.sheet_view.showGridLines = False
db["B1"] = "ניהול תחמושת גדודי — תמונת מצב לחטיבה"
db["B1"].font = F(bold=True, size=16, color=OLIVE)
db["B2"] = 'סה"כ חתום = במכולה (חצרות יסף) + בלבנון (על החיילים) + נצרך (ירי). מזינים אך ורק בגיליון "יומן תנועות" — כל המספרים כאן מחושבים אוטומטית.'
db["B2"].font = F(size=10, italic=True, color=INK2)

L = "'יומן תנועות'"
Q, TY, IT, DT, CO = (f"{L}!$D$2:$D${LOG_MAX}", f"{L}!$A$2:$A${LOG_MAX}",
                     f"{L}!$C$2:$C${LOG_MAX}", f"{L}!$B$2:$B${LOG_MAX}", f"{L}!$E$2:$E${LOG_MAX}")

def sumt(ttype, item_ref):
    return f'SUMIFS({Q},{TY},"{ttype}",{IT},{item_ref})'

# KPI row
kpis = [
    ('סה"כ חתום (יח\')', f'=SUMIFS({Q},{TY},"{T_INTAKE}")', None),
    ("במכולה — חצרות יסף", f'=SUMIFS({Q},{TY},"{T_INTAKE}")-SUMIFS({Q},{TY},"{T_ISSUE}")+SUMIFS({Q},{TY},"{T_RETURN}")', GREEN),
    ("בלבנון — על החיילים", f'=SUMIFS({Q},{TY},"{T_ISSUE}")-SUMIFS({Q},{TY},"{T_RETURN}")-SUMIFS({Q},{TY},"{T_CONSUME}")', BLUE),
    ("נצרך (ירי)", f'=SUMIFS({Q},{TY},"{T_CONSUME}")', ORANGE),
    ("נחתם ב-30 הימים האחרונים", f'=SUMIFS({Q},{TY},"{T_INTAKE}",{DT},">="&TODAY()-30)', None),
]
for k, (lbl, frm, color) in enumerate(kpis):
    col = 2 + k * 2
    lc, vc = db.cell(row=4, column=col), db.cell(row=5, column=col)
    lc.value, vc.value = lbl, frm
    lc.font = F(size=9, color=color or INK2, bold=bool(color))
    vc.font = F(bold=True, size=14)
    vc.number_format = "#,##0"
    lc.alignment = vc.alignment = RIGHT

# item table
HDR_ROW = 7
cols = ["פריט", "קטגוריה", 'סה"כ חתום', "במכולה חצרות יסף", "בלבנון", "נצרך (ירי)", "מאזן"]
for j, h in enumerate(cols, start=2):
    c = db.cell(row=HDR_ROW, column=j, value=h)
    c.font = F(bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=OLIVE)
    c.alignment = CENTER
    c.border = BOX
first, last = HDR_ROW + 1, HDR_ROW + len(ITEMS)
for i, (name, cat) in enumerate(ITEMS):
    r = first + i
    db.cell(row=r, column=2, value=name)
    db.cell(row=r, column=3, value=cat)
    db.cell(row=r, column=4, value=f'={sumt(T_INTAKE, f"$B{r}")}')
    db.cell(row=r, column=5, value=f'={sumt(T_INTAKE, f"$B{r}")}-{sumt(T_ISSUE, f"$B{r}")}+{sumt(T_RETURN, f"$B{r}")}')
    db.cell(row=r, column=6, value=f'={sumt(T_ISSUE, f"$B{r}")}-{sumt(T_RETURN, f"$B{r}")}-{sumt(T_CONSUME, f"$B{r}")}')
    db.cell(row=r, column=7, value=f'={sumt(T_CONSUME, f"$B{r}")}')
    db.cell(row=r, column=8, value=f'=IF(AND(E{r}+F{r}+G{r}=D{r},E{r}>=0,F{r}>=0),"✓","✗")')
tot = last + 1
db.cell(row=tot, column=2, value='סה"כ')
for col in "DEFG":
    db.cell(row=tot, column={"D": 4, "E": 5, "F": 6, "G": 7}[col],
            value=f"=SUM({col}{first}:{col}{last})")
for r in range(first, tot + 1):
    for j in range(2, 9):
        c = db.cell(row=r, column=j)
        c.font = F(size=11, bold=(r == tot or j == 4))
        c.alignment = CENTER if j == 8 else RIGHT
        c.border = BOX
        if j in (4, 5, 6, 7):
            c.number_format = "#,##0"
        if r == tot:
            c.fill = PatternFill("solid", fgColor=OLIVE_SOFT)
db.conditional_formatting.add(
    f"E{first}:F{last}",
    CellIsRule(operator="lessThan", formula=["0"], font=F(color=RED_TXT, bold=True),
               fill=PatternFill("solid", fgColor=RED_FILL)),
)
db.conditional_formatting.add(
    f"H{first}:H{last}",
    CellIsRule(operator="equal", formula=['"✗"'], font=F(color=RED_TXT, bold=True),
               fill=PatternFill("solid", fgColor=RED_FILL)),
)
for j, w in zip(range(2, 9), (20, 16, 13, 17, 11, 12, 8)):
    db.column_dimensions[get_column_letter(j)].width = w
db.freeze_panes = f"A{first}"

# Lebanon-by-company matrix
mstart = tot + 3
db.cell(row=mstart, column=2, value="בלבנון — פירוט לפי פלוגות").font = F(bold=True, size=13, color=OLIVE)
db.cell(row=mstart + 1, column=2,
        value="כמות נוכחית בידי כל פלוגה (ניפוקים פחות החזרות פחות ירי). מיקומי הפלוגות — בגיליון הגדרות.").font = F(size=9, italic=True, color=INK2)
mh = mstart + 2
db.cell(row=mh, column=2, value="פריט").font = F(bold=True, size=10, color="FFFFFF")
db.cell(row=mh, column=2).fill = PatternFill("solid", fgColor=OLIVE)
db.cell(row=mh, column=2).border = BOX
db.cell(row=mh, column=2).alignment = CENTER
for k, (co, locs) in enumerate(COMPANIES):
    c = db.cell(row=mh, column=3 + k, value=co)
    c.font = F(bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=OLIVE)
    c.alignment = CENTER
    c.border = BOX
    db.column_dimensions[get_column_letter(3 + k)].width = 11
for i, (name, _cat) in enumerate(ITEMS):
    r = mh + 1 + i
    db.cell(row=r, column=2, value=name)
    for k in range(len(COMPANIES)):
        col_l = get_column_letter(3 + k)
        db.cell(row=r, column=3 + k, value=(
            f'=SUMIFS({Q},{TY},"{T_ISSUE}",{IT},$B{r},{CO},{col_l}${mh})'
            f'-SUMIFS({Q},{TY},"{T_RETURN}",{IT},$B{r},{CO},{col_l}${mh})'
            f'-SUMIFS({Q},{TY},"{T_CONSUME}",{IT},$B{r},{CO},{col_l}${mh})'))
for r in range(mh + 1, mh + 1 + len(ITEMS)):
    for j in range(2, 3 + len(COMPANIES)):
        c = db.cell(row=r, column=j)
        c.font = F(size=11)
        c.alignment = RIGHT
        c.border = BOX
        if j > 2:
            c.number_format = "#,##0;-#,##0;—"

# ================================================================ הוראות
hp = wb.create_sheet("הוראות")
hp.sheet_view.rightToLeft = True
hp.sheet_view.showGridLines = False
hp.column_dimensions["B"].width = 110
lines = [
    ("איך עובדים עם הקובץ", True),
    ("", False),
    ('1. מזינים אך ורק בגיליון "יומן תנועות" — שורה אחת לכל פעולה. שאר הגיליונות מחושבים אוטומטית ואין לערוך אותם.', False),
    ("2. חתימת מלאי חדש (משיכה מבסיס): סוג תנועה = חתימת מלאי, ממלאים פריט, כמות, תאריך, מקור ושם החותם. הכמות נכנסת למכולה בחצרות יסף.", False),
    ("3. ניפוק ללבנון: סוג תנועה = ניפוק ללבנון, ממלאים פריט, כמות, פלוגה, מיקום ושם מלא של המנפק. הכמות עוברת מהמכולה לפלוגה.", False),
    ("4. דיווח ירי: סוג תנועה = דיווח ירי, עם הפלוגה שירתה. הכמות יורדת מאחזקת הפלוגה ונרשמת כצריכה.", False),
    ("5. החזרה למכולה: סוג תנועה = החזרה למכולה, עם הפלוגה המחזירה.", False),
    ("", False),
    ('פריט/פלוגה/מיקום חדשים — מוסיפים קודם בגיליון "הגדרות" (בשורה פנויה מתחת לרשימה), ואז הם מופיעים ברשימות הנפתחות ביומן.', False),
    ('בדשבורד, עמודת "מאזן" מציגה ✓ כשהמשוואה סגורה. ✗ אדום או מספר שלילי אדום = הוזן ניפוק/ירי מעבר לזמין — בדקו את היומן.', False),
    ("שורה שהוזן בה סוג תנועה בלי פריט או כמות נצבעת באדום ביומן עד להשלמתה.", False),
    ("", False),
    ("שיתוף: העלו את הקובץ ל-Google Drive או OneDrive ושתפו בהרשאת עריכה — כולם עורכים במקביל והדשבורד מתעדכן לבד.", False),
    ("שורות 2–23 ביומן הן יתרות הפתיחה מקובץ המקור (חתימות בילו וחצרות יסף) ומשמשות גם כדוגמת מילוי.", False),
]
for i, (txt, bold) in enumerate(lines, start=2):
    c = hp.cell(row=i, column=2, value=txt)
    c.font = F(bold=bold, size=13 if bold else 11, color=OLIVE if bold else None)
    c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)

wb.save("ammo_management.xlsx")
print("saved ammo_management.xlsx")
