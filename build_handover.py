#!/usr/bin/env python3
"""טופס העברת גזרה והחלפת חתימות - החלפת גדוד מילואים.

גיליונות:
- סקירה: דשבורד - סה"כ פריטים, פריטי צ', ציוד חטיבתי, פערים; פילוח לפי
  קטגוריה ולפי בית/מוצב (מתעדכן אוטומטית מרשימת הציוד).
- רשימת ציוד: הטבלה המרכזית שהגדוד היוצא ממלא - בית, קטגוריה, פריט,
  האם פריט צ' + מספר צ', מקור (חטיבתי/גדודי), כמות ומצב. במעמד ההחלפה
  הגדוד הקולט ממלא 'כמות שנבדקה' ועמודת הפער מחושבת לבד.
- בתים ומוצבים: רשימת הבתים בגזרה (מזינה את הרשימה הנפתחת).
- סיכום וחתימות: פרטי ההעברה, סיכום אוטומטי לפי קטגוריה, הצהרה
  וטבלת חתימות מוסר/קולט/נציג חטיבה - מוכן להדפסה בשני עותקים.
- מקרא והוראות: סדר הפעולות + רשימת הקטגוריות (מזינה את הרשימה הנפתחת).
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

# ---------------------------------------------------------------- נתונים
CATEGORIES = [
    "אמל\"ח",
    "תחמושת",
    "אופטיקה ותצפית",
    "קשר ומחשוב",
    "רחפנים",
    "מיגון",
    "ציוד רפואי",
    "לוגיסטיקה ומגורים",
    "מטבח ומזון",
    "גנרטורים וחשמל",
    "תשתיות ומבנה",
    "כיבוי אש ובטיחות",
    "אחר",
]

SOURCES = ["חטיבתי", "גדודי", "אוגדתי", "אחר"]
CONDITIONS = ["תקין", "תקין חלקית", "תקול", "חסר"]
COND_FILL = {"תקין": "C6EFCE", "תקין חלקית": "FFF2CC", "תקול": "FFC7CE", "חסר": "FFC7CE"}

# שורות דוגמה (למחיקה): בית, קטגוריה, פריט, צ'?, מספר צ', מקור, כמות, מצב
DEMO = [
    ("בית 1", "אופטיקה ותצפית", "משקפת יום", "כן", "100200", "חטיבתי", 2, "תקין"),
    ("בית 1", "קשר ומחשוב", "מכשיר קשר", "כן", "445566", "חטיבתי", 4, "תקין"),
    ("בית 1", "לוגיסטיקה ומגורים", "מיטת שדה + מזרן", "לא", "", "גדודי", 12, "תקין"),
    ("בית 2", "גנרטורים וחשמל", "גנרטור 5 קו\"ט", "כן", "778899", "חטיבתי", 1, "תקול"),
    ("בית 2", "מטבח ומזון", "מקרר", "לא", "", "חטיבתי", 2, "תקין"),
    ("בית 2", "מיגון", "שכפ\"צ", "לא", "", "גדודי", 10, "תקין"),
]

HOUSES_DEMO = ["בית 1", "בית 2", "בית 3"]

# ---------------------------------------------------------------- סגנון
NAVY = "1F4E78"
navy = PatternFill("solid", fgColor=NAVY)
gray = PatternFill("solid", fgColor="D9E1F2")
green = PatternFill("solid", fgColor="C6EFCE")
red = PatternFill("solid", fgColor="FFC7CE")
yellow = PatternFill("solid", fgColor="FFF2CC")
wbf = Font(bold=True, color="FFFFFF", size=11)
bold = Font(bold=True)
big = Font(bold=True, color="FFFFFF", size=15)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center", wrap_text=True)
thin = Side(style="thin", color="9AA7BD")
border = Border(thin, thin, thin, thin)


def title(ws, text, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(1, 1, text)
    c.fill = navy; c.font = big; c.alignment = center
    ws.row_dimensions[1].height = 28


def head(ws, row, labels, widths):
    for c, (lab, w) in enumerate(zip(labels, widths), start=1):
        cell = ws.cell(row, c, lab)
        cell.fill = navy; cell.font = wbf; cell.alignment = center; cell.border = border
        ws.column_dimensions[L(c)].width = w


def rtl(ws):
    ws.sheet_view.rightToLeft = True


def add_dv(ws, formula, cells):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cells)


wb = Workbook()

EQ_SH = "'רשימת ציוד'"
HO_SH = "'בתים ומוצבים'"
LEG_SH = "'מקרא והוראות'"
EROW0, EROW1 = 3, 302        # שורות רשימת הציוד
HROW0, HROW1 = 3, 32         # שורות בתים
CAT_ROW0 = 20                # תחילת רשימת הקטגוריות במקרא
CAT_ROW1 = CAT_ROW0 + len(CATEGORIES) - 1

# ================================================================ סקירה
dash = wb.active
dash.title = "סקירה"
rtl(dash)
title(dash, "העברת גזרה – סיכום ציוד ופערים", 11)

def kpi(r, c, label, formula, fill=None):
    lc = dash.cell(r, c, label); lc.font = bold; lc.alignment = right; lc.border = border
    vc = dash.cell(r, c + 1, formula); vc.font = Font(bold=True, size=12)
    vc.alignment = center; vc.border = border
    if fill:
        vc.fill = fill
    return vc

for c, w in zip(range(1, 12), (20, 9, 3, 18, 9, 9, 9, 9, 3, 14, 9)):
    dash.column_dimensions[L(c)].width = w

dash.cell(3, 1, "סיכום כללי").font = Font(bold=True, size=13)
kpi(4, 1, "סה\"כ שורות ציוד", f"=COUNTA({EQ_SH}!$D${EROW0}:$D${EROW1})", gray)
kpi(5, 1, "סה\"כ כמות נמסרת", f"=SUM({EQ_SH}!$H${EROW0}:$H${EROW1})", gray)
kpi(6, 1, "פריטי צ׳", f'=COUNTIF({EQ_SH}!$E${EROW0}:$E${EROW1},"כן")', yellow)
kpi(7, 1, "ציוד חטיבתי", f'=COUNTIF({EQ_SH}!$G${EROW0}:$G${EROW1},"חטיבתי")')
kpi(8, 1, "ציוד גדודי", f'=COUNTIF({EQ_SH}!$G${EROW0}:$G${EROW1},"גדודי")')

dash.cell(10, 1, "התראות").font = Font(bold=True, size=13)
a1 = kpi(11, 1, "פריטים תקולים",
         f'=COUNTIF({EQ_SH}!$I${EROW0}:$I${EROW1},"תקול")+COUNTIF({EQ_SH}!$I${EROW0}:$I${EROW1},"חסר")')
a2 = kpi(12, 1, "שורות עם פער",
         f'=SUMPRODUCT(({EQ_SH}!$K${EROW0}:$K${EROW1}<>"")*({EQ_SH}!$K${EROW0}:$K${EROW1}<>0))')
a3 = kpi(13, 1, "פריטי צ׳ ללא מספר",
         f'=SUMPRODUCT(({EQ_SH}!$E${EROW0}:$E${EROW1}="כן")*({EQ_SH}!$F${EROW0}:$F${EROW1}=""))')
a4 = kpi(14, 1, "טרם נבדקו ע\"י הקולט",
         f'=SUMPRODUCT(({EQ_SH}!$D${EROW0}:$D${EROW1}<>"")*({EQ_SH}!$J${EROW0}:$J${EROW1}=""))')
for cell in (a1, a2, a3, a4):
    dash.conditional_formatting.add(cell.coordinate,
        CellIsRule(operator="greaterThan", formula=["0"], fill=red, font=bold))

# פילוח לפי קטגוריה
mrow = 3
for j, lab in enumerate(["קטגוריה", "פריטים", "כמות", "פריטי צ׳", "פערים"], start=4):
    c = dash.cell(mrow, j, lab); c.fill = navy; c.font = wbf; c.alignment = center; c.border = border
for i, cat in enumerate(CATEGORIES, start=1):
    r = mrow + i
    rc = dash.cell(r, 4, cat); rc.alignment = right; rc.border = border
    vals = [
        f'=COUNTIF({EQ_SH}!$C${EROW0}:$C${EROW1},$D{r})',
        f'=SUMIF({EQ_SH}!$C${EROW0}:$C${EROW1},$D{r},{EQ_SH}!$H${EROW0}:$H${EROW1})',
        f'=COUNTIFS({EQ_SH}!$C${EROW0}:$C${EROW1},$D{r},{EQ_SH}!$E${EROW0}:$E${EROW1},"כן")',
        f'=SUMPRODUCT(({EQ_SH}!$C${EROW0}:$C${EROW1}=$D{r})*({EQ_SH}!$K${EROW0}:$K${EROW1}<>"")*({EQ_SH}!$K${EROW0}:$K${EROW1}<>0))',
    ]
    for j, v in enumerate(vals, start=5):
        cc = dash.cell(r, j, v); cc.alignment = center; cc.border = border
    dash.conditional_formatting.add(f"H{r}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=red, font=bold))
tr = mrow + len(CATEGORIES) + 1
dash.cell(tr, 4, "סה\"כ").font = bold; dash.cell(tr, 4).alignment = center; dash.cell(tr, 4).border = border
for j in range(5, 9):
    col = L(j)
    cc = dash.cell(tr, j, f"=SUM({col}{mrow+1}:{col}{mrow+len(CATEGORIES)})")
    cc.font = bold; cc.alignment = center; cc.border = border

# פילוח לפי בית/מוצב
hrow = 3
for j, lab in enumerate(["בית / מוצב", "פריטים"], start=10):
    c = dash.cell(hrow, j, lab); c.fill = navy; c.font = wbf; c.alignment = center; c.border = border
for i in range(15):
    r = hrow + 1 + i
    src = f"{HO_SH}!$A${HROW0 + i}"
    dash.cell(r, 10, f'=IF({src}="","",{src})').alignment = right
    dash.cell(r, 11, f'=IF({src}="","",COUNTIF({EQ_SH}!$B${EROW0}:$B${EROW1},{src}))').alignment = center
    dash.cell(r, 10).border = border; dash.cell(r, 11).border = border

# ================================================================ רשימת ציוד
eq = wb.create_sheet("רשימת ציוד")
rtl(eq)
eq_heads = ["מס׳", "בית / מוצב", "קטגוריה", "שם הפריט", "פריט צ׳?", "מספר צ׳",
            "מקור הציוד", "כמות נמסרת", "מצב", "כמות שנבדקה (קולט)", "פער", "הערות"]
title(eq, "רשימת ציוד להעברה – ממולא ע\"י הגדוד המוסר", len(eq_heads))
head(eq, 2, eq_heads, [5, 12, 16, 30, 9, 12, 11, 10, 11, 12, 7, 26])
for i in range(EROW0, EROW1 + 1):
    eq.cell(i, 1, f'=IF($D{i}="","",ROW()-2)').alignment = center
    eq.cell(i, 11, f'=IF(OR($H{i}="",$J{i}=""),"",$H{i}-$J{i})').alignment = center
    for c in range(1, 13):
        cell = eq.cell(i, c)
        cell.border = border
        if cell.alignment.horizontal is None:
            cell.alignment = right if c in (4, 12) else center
for i, (house, cat, item, tz, tznum, src, qty, cond) in enumerate(DEMO):
    r = EROW0 + i
    for c, v in zip((2, 3, 4, 5, 6, 7, 8, 9), (house, cat, item, tz, tznum, src, qty, cond)):
        eq.cell(r, c, v)
    eq.cell(r, 12, "דוגמה – מחק שורה זו")
eq.freeze_panes = "A3"

add_dv(eq, f"={HO_SH}!$A${HROW0}:$A${HROW1}", f"B{EROW0}:B{EROW1}")
add_dv(eq, f"={LEG_SH}!$A${CAT_ROW0}:$A${CAT_ROW1}", f"C{EROW0}:C{EROW1}")
add_dv(eq, '"כן,לא"', f"E{EROW0}:E{EROW1}")
add_dv(eq, '"' + ",".join(SOURCES) + '"', f"G{EROW0}:G{EROW1}")
add_dv(eq, '"' + ",".join(CONDITIONS) + '"', f"I{EROW0}:I{EROW1}")

eq.conditional_formatting.add(f"E{EROW0}:E{EROW1}",
    CellIsRule(operator="equal", formula=['"כן"'], fill=yellow))
eq.conditional_formatting.add(f"F{EROW0}:F{EROW1}",
    FormulaRule(formula=[f'AND($E{EROW0}="כן",$F{EROW0}="")'], fill=red))
for s, color in COND_FILL.items():
    eq.conditional_formatting.add(f"I{EROW0}:I{EROW1}",
        CellIsRule(operator="equal", formula=[f'"{s}"'],
                   fill=PatternFill("solid", fgColor=color)))
eq.conditional_formatting.add(f"K{EROW0}:K{EROW1}",
    FormulaRule(formula=[f'AND($K{EROW0}<>"",$K{EROW0}<>0)'], fill=red, font=bold))
eq.conditional_formatting.add(f"K{EROW0}:K{EROW1}",
    FormulaRule(formula=[f'AND($K{EROW0}<>"",$K{EROW0}=0)'], fill=green))

# ================================================================ בתים ומוצבים
ho = wb.create_sheet("בתים ומוצבים")
rtl(ho)
title(ho, "בתים ומוצבים בגזרה", 4)
head(ho, 2, ["בית / מוצב", "אחראי (גדוד מוסר)", "טלפון", "הערות"], [16, 20, 15, 28])
for i, h in enumerate(HOUSES_DEMO):
    ho.cell(HROW0 + i, 1, h).alignment = center
    ho.cell(HROW0 + i, 4, "דוגמה – ערוך").alignment = right
for r in range(HROW0, HROW1 + 1):
    for c in range(1, 5):
        cell = ho.cell(r, c); cell.border = border
        if cell.alignment.horizontal is None:
            cell.alignment = center if c == 1 else right
ho.freeze_panes = "A3"

# ================================================================ סיכום וחתימות
sig = wb.create_sheet("סיכום וחתימות")
rtl(sig)
SPAN = 6
title(sig, "פרוטוקול העברת גזרה והחלפת חתימות", SPAN)
for c, w in zip(range(1, SPAN + 1), (22, 18, 14, 12, 12, 16)):
    sig.column_dimensions[L(c)].width = w

sig.cell(3, 1, "פרטי ההעברה").font = Font(bold=True, size=13)
details = ["חטיבה", "גזרה", "הגדוד המוסר", "מפקד הגדוד המוסר",
           "הגדוד הקולט", "מפקד הגדוד הקולט", "תאריך ההעברה"]
for i, lab in enumerate(details):
    r = 4 + i
    lc = sig.cell(r, 1, lab); lc.font = bold; lc.alignment = right; lc.border = border
    sig.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    vc = sig.cell(r, 2); vc.fill = gray; vc.alignment = center
    for c in (2, 3):
        sig.cell(r, c).border = border

srow = 4 + len(details) + 1
sig.cell(srow, 1, "סיכום ציוד לפי קטגוריה (אוטומטי)").font = Font(bold=True, size=13)
for j, lab in enumerate(["קטגוריה", "פריטים", "כמות", "פריטי צ׳", "פערים"], start=1):
    c = sig.cell(srow + 1, j, lab); c.fill = navy; c.font = wbf; c.alignment = center; c.border = border
for i, cat in enumerate(CATEGORIES, start=1):
    r = srow + 1 + i
    rc = sig.cell(r, 1, cat); rc.alignment = right; rc.border = border
    vals = [
        f'=COUNTIF({EQ_SH}!$C${EROW0}:$C${EROW1},$A{r})',
        f'=SUMIF({EQ_SH}!$C${EROW0}:$C${EROW1},$A{r},{EQ_SH}!$H${EROW0}:$H${EROW1})',
        f'=COUNTIFS({EQ_SH}!$C${EROW0}:$C${EROW1},$A{r},{EQ_SH}!$E${EROW0}:$E${EROW1},"כן")',
        f'=SUMPRODUCT(({EQ_SH}!$C${EROW0}:$C${EROW1}=$A{r})*({EQ_SH}!$K${EROW0}:$K${EROW1}<>"")*({EQ_SH}!$K${EROW0}:$K${EROW1}<>0))',
    ]
    for j, v in enumerate(vals, start=2):
        cc = sig.cell(r, j, v); cc.alignment = center; cc.border = border
trow = srow + 1 + len(CATEGORIES) + 1
sig.cell(trow, 1, "סה\"כ").font = bold; sig.cell(trow, 1).alignment = center; sig.cell(trow, 1).border = border
for j in range(2, 6):
    col = L(j)
    cc = sig.cell(trow, j, f"=SUM({col}{srow+2}:{col}{srow+1+len(CATEGORIES)})")
    cc.font = bold; cc.alignment = center; cc.border = border

drow = trow + 2
sig.cell(drow, 1, "הצהרה").font = Font(bold=True, size=13)
statements = [
    "אנו החתומים מטה מאשרים כי:",
    "1. הציוד המפורט בגיליון 'רשימת ציוד' נספר ונבדק במעמד נציגי שני הגדודים.",
    "2. הגדוד הקולט מאשר את קבלת הציוד בכמויות ובמצב כמפורט, למעט הפערים שצוינו.",
    "3. פערים שנרשמו יתועדו ויטופלו מול החטיבה.",
    "4. האחריות על הציוד, לרבות פריטי הצ׳, עוברת לגדוד הקולט מרגע החתימה.",
]
for i, t in enumerate(statements):
    r = drow + 1 + i
    sig.merge_cells(start_row=r, start_column=1, end_row=r, end_column=SPAN)
    c = sig.cell(r, 1, t); c.alignment = right

grow = drow + len(statements) + 2
sig.cell(grow, 1, "חתימות").font = Font(bold=True, size=13)
for j, lab in enumerate(["תפקיד", "שם מלא", "דרגה", "מס׳ אישי", "תאריך", "חתימה"], start=1):
    c = sig.cell(grow + 1, j, lab); c.fill = navy; c.font = wbf; c.alignment = center; c.border = border
signers = [
    "הגדוד המוסר – מפקד",
    "הגדוד המוסר – קצין לוגיסטיקה / אפסנאי",
    "הגדוד הקולט – מפקד",
    "הגדוד הקולט – קצין לוגיסטיקה / אפסנאי",
    "נציג חטיבה",
]
for i, role in enumerate(signers):
    r = grow + 2 + i
    sig.row_dimensions[r].height = 24
    rc = sig.cell(r, 1, role); rc.font = bold; rc.alignment = right; rc.border = border
    for c in range(2, SPAN + 1):
        sig.cell(r, c).border = border
last = grow + 2 + len(signers)
sig.merge_cells(start_row=last + 1, start_column=1, end_row=last + 1, end_column=SPAN)
sig.cell(last + 1, 1, "* מומלץ להדפיס בשני עותקים – עותק לכל גדוד.").alignment = right
sig.print_area = f"A1:F{last + 1}"
sig.page_setup.fitToWidth = 1
sig.page_setup.fitToHeight = 1
sig.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================ מקרא והוראות
leg = wb.create_sheet("מקרא והוראות")
rtl(leg)
title(leg, "מקרא והוראות מילוי", 5)
leg.column_dimensions["A"].width = 24
for c in "BCDE":
    leg.column_dimensions[c].width = 18
leg.cell(3, 1, "סדר הפעולות:").font = Font(bold=True, size=13)
steps = [
    "1. הגדוד המוסר ממלא את גיליון 'בתים ומוצבים' – רשימת הבתים/העמדות בגזרה.",
    "2. הגדוד המוסר ממלא את 'רשימת ציוד' – שורה לכל פריט: בית, קטגוריה, שם הפריט,",
    "    האם פריט צ׳ (ואם כן – מספר הצ׳), מקור הציוד (חטיבתי/גדודי), כמות ומצב.",
    "3. במעמד ההחלפה: נציג הגדוד הקולט סופר את הציוד וממלא 'כמות שנבדקה' –",
    "    עמודת הפער מחושבת אוטומטית ונצבעת באדום כשיש אי-התאמה.",
    "4. פריט צ׳ ללא מספר צ׳ נצבע באדום – חובה להשלים לפני חתימה.",
    "5. בודקים את גיליון 'סקירה' – פערים, תקולים ושורות שטרם נבדקו.",
    "6. ממלאים את פרטי ההעברה וחותמים בגיליון 'סיכום וחתימות' (להדפיס בשני עותקים).",
    "",
    "שורות הדוגמה ברשימת הציוד ובבתים מסומנות 'דוגמה' – מחק/ערוך אותן.",
]
for i, t in enumerate(steps):
    r = 4 + i
    leg.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    leg.cell(r, 1, t).alignment = right

leg.cell(15, 1, "צבעים:").font = bold
legend_colors = [(yellow, "פריט צ׳"), (green, "תקין / נבדק ללא פער"),
                 (red, "תקול / חסר / פער / צ׳ בלי מספר")]
for i, (fill, lab) in enumerate(legend_colors):
    sw = leg.cell(16 + i, 1); sw.fill = fill; sw.border = border
    leg.cell(16 + i, 2, lab).alignment = right

cc = leg.cell(CAT_ROW0 - 1, 1, "קטגוריות (מזין את הרשימה הנפתחת):")
cc.font = bold
for i, cat in enumerate(CATEGORIES):
    c = leg.cell(CAT_ROW0 + i, 1, cat); c.alignment = right; c.border = border

OUT = "equipment_handover.xlsx"
wb.save(OUT)
print("נשמר:", OUT, "| גיליונות:", wb.sheetnames)
