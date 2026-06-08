# -*- coding: utf-8 -*-
"""
בונה קובץ ניהול *חי* של טבלת העוצמה הגדודית.
הכול מבוסס נוסחאות וטבלת Excel אמיתית (ListObject) בשם Otzma:
משנים כמות אצל פלוגה / מצאי / מוסיפים פריט — וסה"כ, פער, אחוז מימוש,
סטטוס, הצבעים, הדשבורד, הפילוח והסיכומים מתעדכנים אוטומטית.
"""
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.properties import CalcProperties

SRC = "source_input.xlsx"
OUT = "עוצמה_גדודית_ניהול.xlsx"
TBL = "Otzma"

NAVY="1F3864"; BLUE="2E5496"; STEEL="8EAADB"; RED="C00000"; REDBG="FFC7CE"; REDLT="FCE4E6"
AMBER="BF8F00"; AMBERBG="FFEB9C"; AMBERLT="FFF6DA"; GREEN="375623"; GREENBG="C6EFCE"; GREENLT="EBF4E6"
GREY="808080"; GREYBG="E7E6E6"; GREYLT="F2F2F2"; WHITE="FFFFFF"

thin=Side(style="thin",color="BFBFBF"); med=Side(style="medium",color="808080")
BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
BOX=Border(left=med,right=med,top=med,bottom=med)
def fill(c): return PatternFill("solid", fgColor=c)
CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
CENV=Alignment(horizontal="center",vertical="center")
RGT=Alignment(horizontal="right",vertical="center",wrap_text=True)

# ---------- קטגוריות ----------
CATEGORIES=[
 ("נשק ומקלעים",['רוס"ר M4 פלטופ','מטול+זאבון','פגיון + כוונת','M16','מפרו לייט קצר','זאבון','מקלע נגב קל','נגב 7.62','מאג','מקלרון','רוב"צ HTR 338','רוב"צ M24']),
 ("כוונות ואופטיקה",["כוונת טריג'",'eotech','כוונת השלכה מתקדמת','כוונת השלכה M5','כוונת לילה לצלף MUNS','כוונת טרמית קליפאון','לאופולד','PACK EOTECH','PACK שחור']),
 ('אמר"לים ותצפית',['ישל"ט אורלי גיל','אמר"ל דו עיני מיקרון','אמר"ל דו עיני עדי','אמר"ל דו עיני עידו','אמר"ל שח"ע','אמר"ל שח"מ','אמר"ל אקילה 6*','עמית','חצובה לעמית','עמית איכון','ערמון','ליאור','מכבים','מכבית','מצפן','משקפת','מאתר','חצובה למאתר','שבשבת רוח','מיני תבל נשלט','אמצעי זיהוי תהל']),
 ('סמנים ומציינים',['סמן מפקדים LPL','סמן לייזר למפקד דיויד','ציין לנגב','סמן תרמי נטלי']),
 ('מט"ל ואבזור',['מט"ל מ"מ','מטל ממ VORTEX (אזרחי)','חצובה זיקית (אזרחי)','חצובה מעגל סגור (אזרחי)','חצובה לזום (אזרחי)','ממיר מתח לעמית']),
 ('רחפנים',['רחפן EVOMAX','רחפן AVATA']),
 ('חשמל ואנרגיה',['גנרטור כתום GPT 1800W','גנרטור לבן YAMAR 3800W','גנרטור צהוב 700W','1800W ECOFLOW HYBRID','500W ECOFLOW קטן','1000W ECOFLOW גדול','פלטות','גנרטור דיזל']),
]
CAT_NAMES=[c for c,_ in CATEGORIES]
def category_of(n):
    n=n.strip()
    for cat,items in CATEGORIES:
        for it in items:
            if it.strip()==n: return cat
    return "אחר"

# ---------- קריאת המקור ----------
UNIT_SRC=[("פלוגה א",2),("פלוגה ב",3),("פלוגה ג",4),("מסייעת",5),("פתן",6),("חפק",7),("מפקדה",8)]
src=openpyxl.load_workbook(SRC,data_only=True).active
def num(v):
    if v is None or (isinstance(v,str) and not v.strip()): return None
    try: return int(v) if float(v)%1==0 else float(v)
    except (TypeError,ValueError): return v
records=[]
for r in range(3,src.max_row+1):
    name=src.cell(r,1).value
    if not name or not str(name).strip(): continue
    rec={"name":str(name).strip(),"cat":category_of(str(name).strip())}
    for lbl,c in UNIT_SRC: rec[lbl]=num(src.cell(r,c).value)
    rec["מלאי"]=num(src.cell(r,9).value); rec["מצאי"]=num(src.cell(r,11).value)
    rec["הערה"]=src.cell(r,12).value or ""
    records.append(rec)
# מיון לפי קטגוריה
order=CAT_NAMES+["אחר"]
records.sort(key=lambda r:(order.index(r["cat"]), r["name"]))

# ============================================================
wb=Workbook(); wb.remove(wb.active)
wb.calculation=CalcProperties(fullCalcOnLoad=True)   # חישוב מחדש בכל פתיחה

UNIT_NAMES=[u for u,_ in UNIT_SRC]
HEAD=["קטגוריה","אמצעי"]+UNIT_NAMES+["מלאי","סך הכל","מצאי","פער","אחוז מימוש","סטטוס","הערות","מפתח"]
# מיקומי עמודות (1-based)
COL={h:i+1 for i,h in enumerate(HEAD)}
L=lambda h:get_column_letter(COL[h])
NC=len(HEAD)

# ---------- לשונית רשימות (נסתרת) ----------
lists=wb.create_sheet("רשימות"); lists.sheet_state="hidden"
lists.cell(1,1,"קטגוריות")
for i,c in enumerate(CAT_NAMES+["אחר"]): lists.cell(2+i,1,c)
cat_ref=f"רשימות!$A$2:$A${1+len(CAT_NAMES)+1}"

# ============================================================
# לשונית: טבלת ניהול (חיה)
# ============================================================
ws=wb.create_sheet("טבלת ניהול")
ws.sheet_view.rightToLeft=True
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=NC-1)
t=ws.cell(1,1,"טבלת עוצמה גדודית — גדוד 28  |  קובץ ניהול חי (כל שינוי מתעדכן אוטומטית)")
t.font=Font(bold=True,color=WHITE,size=14); t.fill=fill(NAVY); t.alignment=CENV
ws.row_dimensions[1].height=30

# כותרות (שורה 2)
for h in HEAD:
    c=ws.cell(2,COL[h],h); c.fill=fill(NAVY); c.font=Font(bold=True,color=WHITE,size=10); c.alignment=CEN; c.border=BOX
ws.row_dimensions[2].height=30

first=3; last=first+len(records)-1
for k,rec in enumerate(records):
    r=first+k
    ws.cell(r,COL["קטגוריה"],rec["cat"]).alignment=RGT
    nc=ws.cell(r,COL["אמצעי"],rec["name"]); nc.alignment=RGT; nc.font=Font(bold=True)
    for u in UNIT_NAMES:
        ws.cell(r,COL[u],rec[u]).alignment=CENV
    ws.cell(r,COL["מלאי"],rec["מלאי"]).alignment=CENV
    ws.cell(r,COL["מצאי"],rec["מצאי"]).alignment=CENV
    ws.cell(r,COL["הערות"],rec["הערה"]).alignment=RGT
    # ---- נוסחאות חיות (הפניות יחסיות — תואם כל גרסה, מתמלא אוטומטית בטבלה) ----
    U1=f'{L("פלוגה א")}{r}'; UN=f'{L("מלאי")}{r}'
    K=f'{L("סך הכל")}{r}'; Lm=f'{L("מצאי")}{r}'; M=f'{L("פער")}{r}'; O=f'{L("סטטוס")}{r}'
    ws.cell(r,COL["סך הכל"], f"=SUM({U1}:{UN})").alignment=CENV
    ws.cell(r,COL["פער"], f'=IF({Lm}="","",{K}-{Lm})').alignment=CENV
    pm=ws.cell(r,COL["אחוז מימוש"], f'=IF(OR({Lm}="",{Lm}=0),"",{K}/{Lm})')
    pm.alignment=CENV; pm.number_format="0%"
    ws.cell(r,COL["סטטוס"],
        f'=IF({Lm}="","אין מצאי",IF({K}<{Lm},"חוסר",IF({K}>{Lm},"עודף","תקין")))').alignment=CENV
    ws.cell(r,COL["מפתח"],
        f'=IF({O}="חוסר",{M}-ROW()/100000,IF({O}="עודף",{M}+ROW()/100000,""))')
    for cc in range(1,NC+1):
        ws.cell(r,cc).border=BORDER

# טבלת Excel אמיתית
tab=Table(displayName=TBL, ref=f"A2:{get_column_letter(NC)}{last}")
tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False)
ws.add_table(tab)
ws.freeze_panes="C3"

# רוחב עמודות + הסתרת עמודת מפתח
widths={"קטגוריה":16,"אמצעי":24,"מלאי":8,"סך הכל":8,"מצאי":8,"פער":8,"אחוז מימוש":10,"סטטוס":12,"הערות":22}
for u in UNIT_NAMES: widths[u]=8
for h in HEAD:
    ws.column_dimensions[L(h)].width=widths.get(h,9)
ws.column_dimensions[L("מפתח")].hidden=True

# אימות נתונים: קטגוריה = רשימה נפתחת
dv=DataValidation(type="list",formula1=f"={cat_ref}",allow_blank=True)
dv.add(f"{L('קטגוריה')}{first}:{L('קטגוריה')}{last+200}")  # גם לשורות עתידיות
ws.add_data_validation(dv)

# ---- עיצוב מותנה חי ----
data_rng=f"A{first}:{L('הערות')}{last}"
stat=f"${L('סטטוס')}{first}"
# גוון שורה עדין לפי סטטוס
for txt,col in [("חוסר",REDLT),("עודף",AMBERLT),("תקין",GREENLT),("אין מצאי",GREYLT)]:
    ws.conditional_formatting.add(data_rng,
        FormulaRule(formula=[f'{stat}="{txt}"'], fill=fill(col)))
# עמודת סטטוס — צבע חזק
sr=f"{L('סטטוס')}{first}:{L('סטטוס')}{last}"
for txt,bg,fg in [("חוסר",REDBG,RED),("עודף",AMBERBG,AMBER),("תקין",GREENBG,GREEN),("אין מצאי",GREYBG,GREY)]:
    ws.conditional_formatting.add(sr,
        FormulaRule(formula=[f'{L("סטטוס")}{first}="{txt}"'], fill=fill(bg), font=Font(bold=True,color=fg)))
# עמודת פער — אדום/כתום
pr=f"{L('פער')}{first}:{L('פער')}{last}"
ws.conditional_formatting.add(pr,CellIsRule(operator="lessThan",formula=["0"],fill=fill(REDBG),font=Font(bold=True,color=RED)))
ws.conditional_formatting.add(pr,CellIsRule(operator="greaterThan",formula=["0"],fill=fill(AMBERBG),font=Font(bold=True,color=AMBER)))
# אחוז מימוש — סרגלי נתונים
from openpyxl.formatting.rule import DataBarRule
ws.conditional_formatting.add(f"{L('אחוז מימוש')}{first}:{L('אחוז מימוש')}{last}",
    DataBarRule(start_type="num",start_value=0,end_type="num",end_value=1,color="63BE7B"))

# ============================================================
# לשונית: דשבורד (חי)
# ============================================================
d=wb.create_sheet("דשבורד",0); d.sheet_view.rightToLeft=True; d.sheet_view.showGridLines=False
d.merge_cells("A1:H1")
h=d.cell(1,1,"עוצמה גדודית — לוח בקרה חי"); h.font=Font(bold=True,size=20,color=WHITE); h.fill=fill(NAVY); h.alignment=CENV
d.row_dimensions[1].height=40
d.merge_cells("A2:H2")
s=d.cell(2,1,"כל המספרים מחושבים אוטומטית מלשונית \"טבלת ניהול\" — שנה שם, והכול כאן יתעדכן")
s.font=Font(italic=True,size=10,color="595959"); s.alignment=CENV

S=lambda col: f"{TBL}[{col}]"
cards=[("סך פריטים",f"=COUNTA({S('אמצעי')})",STEEL),
       ("חוסרים",f'=COUNTIF({S("סטטוס")},"חוסר")',RED),
       ("עודפים",f'=COUNTIF({S("סטטוס")},"עודף")',AMBER),
       ("תקין",f'=COUNTIF({S("סטטוס")},"תקין")',GREEN),
       ("ללא מצאי",f'=COUNTIF({S("סטטוס")},"אין מצאי")',GREY),
       ("יח' חסרות",f'=-SUMIF({S("פער")},"<0")',RED),
       ("יח' עודפות",f'=SUMIF({S("פער")},">0")',AMBER)]
col=1
for title,frm,bg in cards:
    d.merge_cells(start_row=4,start_column=col,end_row=4,end_column=col+0)
    d.merge_cells(start_row=5,start_column=col,end_row=6,end_column=col)
    tc=d.cell(4,col,title); tc.fill=fill(bg); tc.font=Font(bold=True,color=WHITE,size=10); tc.alignment=CEN; tc.border=BOX
    vc=d.cell(5,col,frm); vc.fill=fill("F2F2F2"); vc.font=Font(bold=True,size=22,color=bg); vc.alignment=CENV; vc.border=BOX
    d.column_dimensions[get_column_letter(col)].width=14; col+=1
d.row_dimensions[4].height=24; d.row_dimensions[5].height=30; d.row_dimensions[6].height=16

# --- עמודות עזר נסתרות לרשימות החיות (top N) ---
HELP="T"  # מתחיל בעמודה T
hcol=20
# מפתחות חוסרים: SMALL ; עודפים: LARGE
def help_block(start_row, kind, n):
    for k in range(1,n+1):
        rr=start_row+k-1
        if kind=="short":
            keyf=f'=IFERROR(SMALL({S("מפתח")},{k}),"")'
        else:
            keyf=f'=IFERROR(IF(LARGE({S("מפתח")},{k})>0,LARGE({S("מפתח")},{k}),""),"")'
        d.cell(rr,hcol,keyf)
    return
help_block(2,"short",12)
help_block(20,"surp",6)
for c in range(hcol,hcol+1):
    d.column_dimensions[get_column_letter(c)].hidden=True

def live_list(r0, title, headcol, key_start, n, posneg):
    d.merge_cells(start_row=r0,start_column=1,end_row=r0,end_column=4)
    hc=d.cell(r0,1,title); hc.fill=fill(headcol); hc.font=Font(bold=True,color=WHITE,size=12)
    hc.alignment=Alignment(horizontal="right",vertical="center",indent=1)
    d.row_dimensions[r0].height=24
    for i,htxt in enumerate(["אמצעי","סך הכל","מצאי","פער"]):
        c=d.cell(r0+1,1+i,htxt); c.fill=fill(NAVY); c.font=Font(bold=True,color=WHITE,size=10); c.alignment=CEN; c.border=BOX
    for k in range(n):
        rr=r0+2+k; keycell=f"${get_column_letter(hcol)}${key_start+k}"
        mt=f'MATCH({keycell},{S("מפתח")},0)'
        d.cell(rr,1,f'=IFERROR(INDEX({S("אמצעי")},{mt}),"")').alignment=RGT
        d.cell(rr,2,f'=IFERROR(INDEX({S("סך הכל")},{mt}),"")').alignment=CENV
        d.cell(rr,3,f'=IFERROR(INDEX({S("מצאי")},{mt}),"")').alignment=CENV
        gc=d.cell(rr,4,f'=IFERROR(INDEX({S("פער")},{mt}),"")'); gc.alignment=CENV
        gc.font=Font(bold=True,color=(RED if posneg<0 else AMBER))
        gc.fill=fill(REDBG if posneg<0 else AMBERBG)
        for cc in range(1,5): d.cell(rr,cc).border=BORDER
live_list(8,"חוסרים מובילים (מתעדכן)",RED,2,12,-1)

# בנה את העודפים בצד שמאל (עמודות 5-8)
for c in range(5,9): d.column_dimensions[get_column_letter(c)].width=14
def live_list_right(r0,title,headcol,key_start,n,posneg):
    d.merge_cells(start_row=r0,start_column=5,end_row=r0,end_column=8)
    hc=d.cell(r0,5,title); hc.fill=fill(headcol); hc.font=Font(bold=True,color=WHITE,size=12)
    hc.alignment=Alignment(horizontal="right",vertical="center",indent=1);
    for i,htxt in enumerate(["אמצעי","סך הכל","מצאי","פער"]):
        c=d.cell(r0+1,5+i,htxt); c.fill=fill(NAVY); c.font=Font(bold=True,color=WHITE,size=10); c.alignment=CEN; c.border=BOX
    for k in range(n):
        rr=r0+2+k; keycell=f"${get_column_letter(hcol)}${key_start+k}"
        mt=f'MATCH({keycell},{S("מפתח")},0)'
        d.cell(rr,5,f'=IFERROR(INDEX({S("אמצעי")},{mt}),"")').alignment=RGT
        d.cell(rr,6,f'=IFERROR(INDEX({S("סך הכל")},{mt}),"")').alignment=CENV
        d.cell(rr,7,f'=IFERROR(INDEX({S("מצאי")},{mt}),"")').alignment=CENV
        gc=d.cell(rr,8,f'=IFERROR(INDEX({S("פער")},{mt}),"")'); gc.alignment=CENV
        gc.font=Font(bold=True,color=AMBER); gc.fill=fill(AMBERBG)
        for cc in range(5,9): d.cell(rr,cc).border=BORDER
live_list_right(8,"עודפים מובילים (מתעדכן)",AMBER,20,6,1)
for c in range(1,5): d.column_dimensions[get_column_letter(c)].width=15

# ============================================================
# לשונית: סיכום לפי קטגוריה (חי)
# ============================================================
cs=wb.create_sheet("סיכום קטגוריה"); cs.sheet_view.rightToLeft=True
cols=["קטגוריה","פריטים","סך יחידות","סך מצאי","סך פער","חוסרים","עודפים","תקין"]
cs.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
tt=cs.cell(1,1,"סיכום לפי קטגוריה (חי)"); tt.fill=fill(NAVY); tt.font=Font(bold=True,color=WHITE,size=13); tt.alignment=CENV
cs.row_dimensions[1].height=28
for i,htxt in enumerate(cols):
    c=cs.cell(2,1+i,htxt); c.fill=fill(BLUE); c.font=Font(bold=True,color=WHITE); c.alignment=CEN; c.border=BOX
rr=3
for cat in CAT_NAMES+["אחר"]:
    cc=f'"{cat}"'
    cs.cell(rr,1,cat).alignment=RGT; cs.cell(rr,1).font=Font(bold=True)
    cs.cell(rr,2,f'=COUNTIF({S("קטגוריה")},{cc})').alignment=CENV
    cs.cell(rr,3,f'=SUMIF({S("קטגוריה")},{cc},{S("סך הכל")})').alignment=CENV
    cs.cell(rr,4,f'=SUMIF({S("קטגוריה")},{cc},{S("מצאי")})').alignment=CENV
    gc=cs.cell(rr,5,f'=SUMIFS({S("פער")},{S("קטגוריה")},{cc})'); gc.alignment=CENV
    cs.cell(rr,6,f'=COUNTIFS({S("קטגוריה")},{cc},{S("סטטוס")},"חוסר")').alignment=CENV
    cs.cell(rr,7,f'=COUNTIFS({S("קטגוריה")},{cc},{S("סטטוס")},"עודף")').alignment=CENV
    cs.cell(rr,8,f'=COUNTIFS({S("קטגוריה")},{cc},{S("סטטוס")},"תקין")').alignment=CENV
    for c in range(1,len(cols)+1): cs.cell(rr,c).border=BORDER
    cs.conditional_formatting.add(f"E{rr}",CellIsRule(operator="lessThan",formula=["0"],fill=fill(REDBG),font=Font(bold=True,color=RED)))
    cs.conditional_formatting.add(f"E{rr}",CellIsRule(operator="greaterThan",formula=["0"],fill=fill(AMBERBG),font=Font(bold=True,color=AMBER)))
    rr+=1
# שורת סך-הכל
cs.cell(rr,1,"סה\"כ").font=Font(bold=True,color=WHITE); cs.cell(rr,1).fill=fill(NAVY); cs.cell(rr,1).alignment=RGT
for i,frm in enumerate([f'=SUM(B3:B{rr-1})',f'=SUM(C3:C{rr-1})',f'=SUM(D3:D{rr-1})',f'=SUM(E3:E{rr-1})',f'=SUM(F3:F{rr-1})',f'=SUM(G3:G{rr-1})',f'=SUM(H3:H{rr-1})']):
    c=cs.cell(rr,2+i,frm); c.fill=fill(NAVY); c.font=Font(bold=True,color=WHITE); c.alignment=CENV
for i,w in enumerate([20,10,12,10,10,10,10,10]): cs.column_dimensions[get_column_letter(i+1)].width=w

# ============================================================
# לשונית: פילוח לפי פלוגה (מטריצה חיה SUMIFS)
# ============================================================
pf=wb.create_sheet("פילוח פלוגה"); pf.sheet_view.rightToLeft=True
pcols=["קטגוריה"]+UNIT_NAMES+["מלאי","סך הכל"]
pf.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(pcols))
tt=pf.cell(1,1,"פילוח אחזקה לפי פלוגה / גורם (חי)"); tt.fill=fill(NAVY); tt.font=Font(bold=True,color=WHITE,size=13); tt.alignment=CENV
pf.row_dimensions[1].height=28
for i,htxt in enumerate(pcols):
    c=pf.cell(2,1+i,htxt); c.fill=fill(BLUE); c.font=Font(bold=True,color=WHITE); c.alignment=CEN; c.border=BOX
rr=3
sum_cols=UNIT_NAMES+["מלאי","סך הכל"]
for cat in CAT_NAMES+["אחר"]:
    cc=f'"{cat}"'
    pf.cell(rr,1,cat).alignment=RGT; pf.cell(rr,1).font=Font(bold=True)
    for j,sc in enumerate(sum_cols):
        pf.cell(rr,2+j,f'=SUMIF({S("קטגוריה")},{cc},{S(sc)})').alignment=CENV
    for c in range(1,len(pcols)+1): pf.cell(rr,c).border=BORDER
    rr+=1
pf.cell(rr,1,"סה\"כ").font=Font(bold=True,color=WHITE); pf.cell(rr,1).fill=fill(NAVY); pf.cell(rr,1).alignment=RGT
for j,sc in enumerate(sum_cols):
    cl=get_column_letter(2+j); c=pf.cell(rr,2+j,f'=SUM({cl}3:{cl}{rr-1})')
    c.fill=fill(NAVY); c.font=Font(bold=True,color=WHITE); c.alignment=CENV
for i,w in enumerate([16]+[9]*len(UNIT_NAMES)+[9,9]): pf.column_dimensions[get_column_letter(i+1)].width=w

# ============================================================
# לשונית: מקרא והוראות
# ============================================================
lg=wb.create_sheet("מקרא והוראות"); lg.sheet_view.rightToLeft=True; lg.sheet_view.showGridLines=False
lg.merge_cells("A1:B1")
tt=lg.cell(1,1,"איך עובדים עם הקובץ"); tt.fill=fill(NAVY); tt.font=Font(bold=True,color=WHITE,size=14); tt.alignment=CENV
lg.row_dimensions[1].height=32
rows=[
 ("עיקרון","הקובץ חי: עורכים רק בלשונית \"טבלת ניהול\" — כל השאר מחושב לבד",STEEL),
 ("שינוי כמות","הקלד/שנה מספר אצל פלוגה, מלאי או מצאי — סך הכל, פער, אחוז המימוש, הסטטוס והצבע יתעדכנו מיד",None),
 ("הוספת פריט","כתוב בשורה הריקה מתחת לטבלה — הטבלה תתרחב והנוסחאות יושלמו אוטומטית. בחר קטגוריה מהרשימה הנפתחת",None),
 ("מחיקת פריט","מחק את כל השורה (קליק ימני> מחק שורת טבלה)",None),
 ("סך הכל","סכום הפלוגות + מלאי (נוסחה)",None),
 ("פער","סך הכל פחות מצאי",None),
 ("אחוז מימוש","סך הכל חלקי מצאי (סרגל נתונים)",None),
 ("חוסר (אדום)","סך הכל קטן מהמצאי הרשום",REDBG),
 ("עודף (כתום)","סך הכל גדול מהמצאי הרשום",AMBERBG),
 ("תקין (ירוק)","התאמה מלאה",GREENBG),
 ("אין מצאי (אפור)","לא הוזן ערך מצאי",GREYBG),
 ("סינון","בלשונית הטבלה אפשר ללחוץ על חץ הסינון בכותרת ולהציג למשל רק \"חוסר\"",None),
]
r=3
for a,b,col in rows:
    ca=lg.cell(r,1,a); ca.font=Font(bold=True); ca.alignment=RGT; ca.border=BORDER
    cb=lg.cell(r,2,b); cb.alignment=RGT; cb.border=BORDER
    if col: ca.fill=fill(col)
    r+=1
lg.column_dimensions["A"].width=20; lg.column_dimensions["B"].width=72

# סדר לשוניות
tab_order=["דשבורד","טבלת ניהול","סיכום קטגוריה","פילוח פלוגה","מקרא והוראות","רשימות"]
wb._sheets.sort(key=lambda s: tab_order.index(s.title) if s.title in tab_order else 99)
wb.active=0
wb.save(OUT)
print("נשמר:",OUT,"| פריטים:",len(records))
