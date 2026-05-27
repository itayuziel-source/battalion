#!/usr/bin/env python3
"""תוכנית חיתוך פח עם ערבוב סוגי פלטות (ללא חפיפות).

לכל סוג פאנל נבחרת הפלטה הכי משתלמת (הכי פחות שטח/פחת), הפאנל נחתך
לחתיכות בהתאם, וכל החתיכות של אותו סוג פלטה נארזות יחד (ניצול שאריות)
באמצעות MaxRects עם סיבוב. הפלט: לכל פלטה - אילו חתיכות ומאיזה מקור.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from sheet_calc import PANELS, STOCKS, split_panel, MaxRectsBin

REQ_AREA = sum(w * h * q for w, h, q in PANELS)


def pack_record(items, W, H):
    """items: (w,h,label). מחזיר רשימת פלטות עם מיקומי חתיכות."""
    items = sorted(items, key=lambda p: -(p[0] * p[1]))
    plates, objs = [], []
    for (w, h, label) in items:
        placed = False
        for i, bo in enumerate(objs):
            pos = bo.insert(w, h)
            if pos:
                plates[i]["pieces"].append((label,) + pos)
                placed = True
                break
        if not placed:
            bo = MaxRectsBin(W, H)
            pos = bo.insert(w, h)
            objs.append(bo)
            plates.append({"plate": (W, H), "pieces": [(label,) + pos]})
    return plates


def best_plate(w, h, q):
    """בוחר את הפלטה עם הכי מעט שטח כולל לפאנל הזה."""
    best = None
    for (pw, ph) in STOCKS:
        items = [(a, b, "") for (a, b) in split_panel(w, h, pw, ph)] * q
        n = len(pack_record(items, pw, ph))
        area = n * pw * ph
        if best is None or area < best[0]:
            best = (area, (pw, ph))
    return best[1]


# שיוך פאנלים לפלטות וקיבוץ
pools = {s: [] for s in STOCKS}
assignment = []
for (w, h, q) in PANELS:
    pw, ph = best_plate(w, h, q)
    assignment.append((w, h, q, (pw, ph)))
    sub = split_panel(w, h, pw, ph)
    for inst in range(1, q + 1):
        for si, (a, b) in enumerate(sub, start=1):
            lbl = f"{w}×{h} #{inst}"
            if len(sub) > 1:
                lbl += f" ח{si}/{len(sub)}"
            pools[(pw, ph)].append((a, b, lbl))

# אריזה לכל סוג פלטה
all_plates = []
for s, items in pools.items():
    if items:
        all_plates.extend(pack_record(items, *s))

total_area = sum(p["plate"][0] * p["plate"][1] for p in all_plates)
waste = 1 - REQ_AREA / total_area
mix = {}
for p in all_plates:
    mix[p["plate"]] = mix.get(p["plate"], 0) + 1

print(f"שטח נדרש: {REQ_AREA/1e6:.2f} מ״ר | סה״כ פלטות: {len(all_plates)} | "
      f"שטח כולל: {total_area/1e6:.2f} מ״ר | פחת: {waste*100:.0f}%")
for s, c in sorted(mix.items()):
    print(f"  {s[0]}x{s[1]}: {c} פלטות")

# ---------------------------------------------------------------- אקסל
NAVY = "1F4E78"
fill_navy = PatternFill("solid", fgColor=NAVY)
fill_hdr = PatternFill("solid", fgColor="2E75B6")
fill_gray = PatternFill("solid", fgColor="D9E1F2")
fill_green = PatternFill("solid", fgColor="C6EFCE")
white_bold = Font(bold=True, color="FFFFFF", size=11)
bold = Font(bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")
thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# --- תוכנית חיתוך
cut = wb.active
cut.title = "תוכנית חיתוך"
cut.sheet_view.rightToLeft = True
heads = ["לוח #", "מידת פלטה (מ״מ)", "חתיכה (מקור)",
         "רוחב חתיכה", "אורך חתיכה", "X (מ״מ)", "Y (מ״מ)"]
widths = [8, 16, 22, 12, 12, 10, 10]
for c, (lab, w) in enumerate(zip(heads, widths), start=1):
    cell = cut.cell(1, c, lab)
    cell.fill = fill_navy; cell.font = white_bold; cell.alignment = center; cell.border = border
    cut.column_dimensions[L(c)].width = w

row = 2
for idx, p in enumerate(all_plates, start=1):
    pw, ph = p["plate"]
    used = sum(pc[3] * pc[4] for pc in p["pieces"])
    util = used / (pw * ph)
    hc = cut.cell(row, 1, f"לוח {idx}")
    cut.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    hc.fill = fill_hdr; hc.font = white_bold; hc.alignment = center
    info = cut.cell(row, 3, f"{pw}x{ph} | ניצול {util*100:.0f}% | {len(p['pieces'])} חתיכות")
    cut.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
    info.fill = fill_hdr; info.font = white_bold; info.alignment = center
    row += 1
    for (label, x, y, w_, h_) in sorted(p["pieces"], key=lambda z: (z[2], z[1])):
        vals = [idx, f"{pw}x{ph}", label, w_, h_, x, y]
        for c, v in enumerate(vals, start=1):
            cell = cut.cell(row, c, v); cell.alignment = center; cell.border = border
        row += 1
cut.freeze_panes = "A2"

# --- סיכום
sm = wb.create_sheet("סיכום")
sm.sheet_view.rightToLeft = True
sm.merge_cells("A1:C1")
t = sm.cell(1, 1, "סיכום תוכנית חיתוך - ערבוב פלטות")
t.fill = fill_navy; t.font = Font(bold=True, color="FFFFFF", size=14); t.alignment = center
for c, lab in enumerate(["מידת פלטה (מ״מ)", "כמות פלטות", "שטח (מ״ר)"], start=1):
    cell = sm.cell(2, c, lab); cell.fill = fill_navy; cell.font = white_bold
    cell.alignment = center; cell.border = border
    sm.column_dimensions[L(c)].width = 18
r2 = 3
for s, c in sorted(mix.items()):
    sm.cell(r2, 1, f"{s[0]}x{s[1]}").alignment = center
    sm.cell(r2, 2, c).alignment = center
    sm.cell(r2, 3, round(c * s[0] * s[1] / 1e6, 2)).alignment = center
    for cc in range(1, 4):
        sm.cell(r2, cc).border = border
    r2 += 1
sm.cell(r2, 1, "סה״כ").font = bold; sm.cell(r2, 1).fill = fill_gray; sm.cell(r2, 1).alignment = center
sm.cell(r2, 2, len(all_plates)).font = bold; sm.cell(r2, 2).fill = fill_gray; sm.cell(r2, 2).alignment = center
sm.cell(r2, 3, round(total_area / 1e6, 2)).font = bold; sm.cell(r2, 3).fill = fill_gray; sm.cell(r2, 3).alignment = center
sm.cell(r2 + 2, 1, "שטח נדרש נטו (מ״ר):").alignment = right; sm.cell(r2 + 2, 1).font = bold
sm.cell(r2 + 2, 2, round(REQ_AREA / 1e6, 2)).alignment = center
wc = sm.cell(r2 + 3, 1, "פחת כולל:"); wc.alignment = right; wc.font = bold
wv = sm.cell(r2 + 3, 2, round(waste, 4)); wv.number_format = "0%"; wv.alignment = center; wv.fill = fill_green

# --- שיוך פאנל->פלטה
asg = wb.create_sheet("שיוך פאנלים")
asg.sheet_view.rightToLeft = True
for c, lab in enumerate(["פאנל (מ״מ)", "כמות", "פלטה נבחרת", "חתיכות לפאנל"], start=1):
    cell = asg.cell(1, c, lab); cell.fill = fill_navy; cell.font = white_bold
    cell.alignment = center; cell.border = border
    asg.column_dimensions[L(c)].width = 16
for i, (w, h, q, (pw, ph)) in enumerate(assignment, start=2):
    n_pieces = len(split_panel(w, h, pw, ph))
    for c, v in enumerate([f"{w}x{h}", q, f"{pw}x{ph}", n_pieces], start=1):
        cell = asg.cell(i, c, v); cell.alignment = center; cell.border = border

OUT = "cut_plan.xlsx"
wb.save(OUT)
print(f"נשמר: {OUT}")
