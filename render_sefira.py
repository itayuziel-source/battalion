# -*- coding: utf-8 -*-
"""תצוגה מקדימה של בלוק הספירה המתוקן (מחשב בפייתון מה ש-Excel יחשב)."""
import openpyxl, matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from matplotlib import font_manager as fm
from bidi.algorithm import get_display
for f in ["/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
          "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"]:
    fm.fontManager.addfont(f)
plt.rcParams["font.family"]="Liberation Sans"
def H(s): return get_display("" if s is None else str(s))
NAVY="#1F3864"; RED="#C00000"; REDBG="#FFC7CE"; GREEN="#375623"; HDR="#2E5496"; LBL="#D9E1F2"

wb=openpyxl.load_workbook("sefira_input.xlsx", data_only=True)
ws=wb["שבצק יציאות"]
TS='ת״ש'
S=range(5,30)
roles={r: (ws.cell(r,2).value or "") for r in S}
names={r: (ws.cell(r,1).value or "") for r in S}
n_soldiers=sum(1 for r in S if str(names[r]).strip())
DAYS=list(range(4,106))  # D..DA
NDAY=21  # נציג 3 שבועות ראשונים
mins={"base":ws["B39"].value,"מחתים":ws["B40"].value,"לוגיסטיקה":ws["B41"].value}

def day_hdr(c):
    h=ws.cell(3,c).value or ""
    parts=str(h).split("\n"); return parts[0], (parts[1] if len(parts)>1 else "")

rows_def=[
 ("סה\"כ בבסיס","base",True),
 ("מחתים בבסיס","מחתים",True),
 ("לוגיסטיקה בבסיס","לוגיסטיקה",True),
 ("סה\"כ בבית","בית",False),
 ("סה\"כ לא בבסיס","away",False),
 ("__sep__","",False),
 ("חופשה","חופשה",False),(TS,TS,False),("קורס","קורס",False),
 ("מבחן","מבחן",False),("חול","חול",False),("בית","בית",False),
 ("X","X",False),("? לא ידוע","?",False),("ללא סטטוס","ריק",False),
]
def val(c,kind):
    cells=[ws.cell(r,c).value for r in S]
    cells=[("" if v is None else str(v).strip()) for v in cells]
    if kind=="base": return sum(1 for v in cells if v=="בסיס")
    if kind=="away": return sum(1 for v in cells if v!="" and v!="בסיס")
    if kind=="ריק":  return n_soldiers-sum(1 for v in cells if v!="")
    if kind in ("מחתים","לוגיסטיקה"):
        return sum(1 for r in S if kind in str(roles[r]) and (str(ws.cell(r,c).value).strip() if ws.cell(r,c).value else "")=="בסיס")
    return sum(1 for v in cells if v==kind)

ncols=NDAY+1
fig_w=ncols*0.52+3.2; fig_h=len(rows_def)*0.42+1.4
fig,ax=plt.subplots(figsize=(fig_w,fig_h)); ax.axis("off")
W=0.52; LBLW=3.2; rh=0.42
total_w=LBLW+NDAY*W
ax.set_xlim(0,total_w); ax.set_ylim(0,(len(rows_def)+2)*rh); ax.invert_yaxis()
# כותרת
ax.add_patch(Rectangle((0,0),total_w,rh*1.2,color=NAVY))
ax.text(total_w-0.1,rh*0.6,H("בלוק הספירה המתוקן — מי בבסיס ומי לא (21/6–11/7 לדוגמה)"),color="white",fontsize=11,fontweight="bold",ha="right",va="center")
y=rh*1.2
# שורת תאריכים
ax.add_patch(Rectangle((total_w-LBLW,y),LBLW,rh,color=HDR))
ax.text(total_w-0.1,y+rh/2,H("יום ◂"),color="white",fontsize=8,fontweight="bold",ha="right",va="center")
for j,c in enumerate(DAYS[:NDAY]):
    d,wd=day_hdr(c); x=total_w-LBLW-(j+1)*W
    shab = wd=="ש"
    ax.add_patch(Rectangle((x,y),W,rh,facecolor="#8EAADB" if shab else HDR,edgecolor="white",lw=0.4))
    ax.text(x+W/2,y+rh/2,H(f"{d}\n{wd}"),color="white",fontsize=6.5,ha="center",va="center",fontweight="bold")
y+=rh
for label,kind,is_min in rows_def:
    if label=="__sep__": y+=rh*0.3; continue
    ax.add_patch(Rectangle((total_w-LBLW,y),LBLW,rh,facecolor=LBL,edgecolor="#bbb",lw=0.4))
    ax.text(total_w-0.1,y+rh/2,H(label),fontsize=7.5,fontweight="bold",ha="right",va="center")
    for j,c in enumerate(DAYS[:NDAY]):
        v=val(c,kind); x=total_w-LBLW-(j+1)*W
        bg="white"; fg="#222"
        if is_min and v<mins[kind]: bg=REDBG; fg=RED
        ax.add_patch(Rectangle((x,y),W,rh,facecolor=bg,edgecolor="#ddd",lw=0.3))
        ax.text(x+W/2,y+rh/2,H(v),fontsize=7,ha="center",va="center",
                fontweight="bold" if (is_min) else "normal",color=fg)
    y+=rh
plt.tight_layout(); plt.savefig("preview_sefira_count.png",dpi=140,bbox_inches="tight"); plt.close()
print("saved. soldiers:",n_soldiers,"| mins:",mins)
