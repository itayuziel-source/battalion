#!/usr/bin/env python3
"""מערכת ניהול רחפנים - ארמו"ן גדוד 28.

גיליונות מקושרים (הצלבות):
- סקירה: דשבורד עם ספירות והתראות (COUNTIF/COUNTIFS/SUMPRODUCT על הצי).
- צי הרחפנים: רישום כל רחפן (צ', סוג, ערכה, סטטוס, מחזיק, תאריכים, מצב).
  עמודת 'פריטי צ׳ נדרשים' נשלפת אוטומטית מטבלת הסוגים (VLOOKUP).
- יומן השאלות: תיעוד השאלה/החזרה; סוג הרחפן נשלף לפי הצ' (VLOOKUP).
- תכולת ערכת EVO / AVATA: רשימת פריטים + סימון 'צ׳' + ספירת חוסרים.
- חיילים: רשימה למילוי (מקור לרשימות נפתחות של מחזיקים).
- מקרא: סטטוסים, מצבים, וטבלת סוגים (סופרת פריטים מהתכולות).
"""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

# ---------------------------------------------------------------- תכולות
EVO_ITEMS = [
    ("רחפן EVO + מגן גימבל + 4 פרופים 1136/1158", "כן"),
    ("שלט לרחפן + 2 אנטנות + 2 סטיקים", "כן"),
    ("סוללה לרחפן", "לא"),
    ('שנאי סוללה + כבל "שמינייה"', "לא"),
    ("ראש מטען קיר 65W", "לא"),
    ("כבל USB – TYPE C", "לא"),
    ("כבל TYPE C – TYPE C", "לא"),
    ("פרופים ספייר סדרה 1136/1158", "לא"),
    ('תיק פקל EVO ירוק 5 חלקים (תא מרכזי, תא צד סוללות, תא צד אחסון, רצועה, מחיצה)', "כן"),
    ("התקן הטלה כדור ברזל", "כן"),
    ("התקן הטלה כדור נוצה", "כן"),
    ('מטען רב ערוצי + כבל "שמינייה"', "לא"),
    ("אנטנה מגדיל טווח + זוג מתאמים לשלט", "כן"),
    ("חצובה לאנטנה מגדיל טווח", "לא"),
    ("כבל מאריך 20 מטר לאנטנה מגדיל טווח", "לא"),
    ("ערכת טעינה לשטח (תיק, כבלי 12V, מטען רכב 200/240W, כבלי C-C, USB-C 7A, מתאם עם צג מתח)", "כן"),
]

AVATA_ITEMS = [
    ("תיק ירך 2 חלקים – רצועות חגורה + רצועת נשיאה + קייס למשקף + פאוץ' צג מפקד", "כן"),
    ("רחפן AVATA + מגן גימבל + 4 פרופים", "כן"),
    ("ג'ויסטיק + רצועה", "כן"),
    ("משקף אינטגרה", "כן"),
    ("משקף גוגלס 2 – רצועת ראש + מגן עדשות + כבל סוללה + סוללה למשקף", "כן"),
    ("כרטיס זיכרון למשקף 128GB אקסטרים", "לא"),
    ("סוללה לרחפן", "לא"),
    ("מטען סוללת רחפן יחיד", "לא"),
    ("רכזת טעינה (4 סוללות) לרחפן", "לא"),
    ("ראש מטען קיר", "לא"),
    ("כבל USB – TYPE C", "לא"),
    ("כבל TYPE C – TYPE C", "לא"),
    ("צג מפקד (סמארטפון)", "כן"),
    ("זוג פרופים ספייר", "לא"),
    ("ערכת תאורה A – 4 פנסים, תושבת מרכזית כפולה, תושבת L, תושבת R", "לא"),
    ("ערכת תאורה B – 2 פנסים כפולים, תושבת מרכזית + תושבת צד", "לא"),
]

SOLDIERS = ["דוד כהן", "מאור לוי", "יוסי בר", "איתי שמש", "ניב סורוקר",
            "אורי עוזיאל", "גבריאל יעקוביאן", "עומר נוימן"]

# צי לדוגמה (מחק/ערוך): צ', סוג, ערכה, סטטוס, מחזיק, השאלה, יעד, מצב, הערות
FLEET = [
    ("EVO-01", "EVO", "1136", "הושאל", "דוד כהן", "2026-05-20", "2026-06-20", "תקין", ""),
    ("EVO-02", "EVO", "1158", "במחסן", "", "", "", "תקין", ""),
    ("EVO-03", "EVO", "1136", "בתיקון", "", "", "", "תקול", "גימבל לא מתכייל"),
    ("EVO-04", "EVO", "1158", "מושבת", "", "", "", "תקול", "ממתין לחלקים"),
    ("AVATA-01", "AVATA", "1008", "הושאל", "מאור לוי", "2026-04-01", "2026-05-01", "תקין", ""),
    ("AVATA-02", "AVATA", "1008", "במחסן", "", "", "", "תקין", ""),
    ("AVATA-03", "AVATA", "1008", "הושאל", "יוסי בר", "2026-05-25", "2026-06-25", "תקין", ""),
    ("AVATA-04", "AVATA", "1008", "במחסן", "", "", "", "תקין", ""),
]

LOG = [
    ("2026-05-20", "EVO-01", "השאלה", "דוד כהן", 'סמל אפסנאות', ""),
    ("2026-04-01", "AVATA-01", "השאלה", "מאור לוי", 'סמל אפסנאות', ""),
    ("2026-05-25", "AVATA-03", "השאלה", "יוסי בר", 'סמל אפסנאות', ""),
]

STATUSES = ["במחסן", "הושאל", "בתיקון", "מושבת"]
STATUS_FILL = {"במחסן": "C6EFCE", "הושאל": "FFF2CC", "בתיקון": "FFD966", "מושבת": "FFC7CE"}

# ---------------------------------------------------------------- סגנון
NAVY = "1F4E78"
navy = PatternFill("solid", fgColor=NAVY)
sub = PatternFill("solid", fgColor="2E75B6")
gray = PatternFill("solid", fgColor="D9E1F2")
green = PatternFill("solid", fgColor="C6EFCE")
red = PatternFill("solid", fgColor="FFC7CE")
yellow = PatternFill("solid", fgColor="FFF2CC")
wbf = Font(bold=True, color="FFFFFF", size=11)
bold = Font(bold=True)
normal = Font(size=11)
big = Font(bold=True, color="FFFFFF", size=15)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center", wrap_text=True)
thin = Side(style="thin", color="9AA7BD")
border = Border(thin, thin, thin, thin)


def title(ws, text, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(1, 1, text)
    c.fill = navy; c.font = big; c.alignment = center
    ws.row_dimensions[1].height = 26


def head(ws, row, labels, widths):
    for c, (lab, w) in enumerate(zip(labels, widths), start=1):
        cell = ws.cell(row, c, lab)
        cell.fill = navy; cell.font = wbf; cell.alignment = center; cell.border = border
        ws.column_dimensions[L(c)].width = w


def rtl(ws):
    ws.sheet_view.rightToLeft = True


def D(s):
    return datetime.strptime(s, "%Y-%m-%d").date() if s else ""


wb = Workbook()

# רפרנסים לטווחים (לשמות גיליונות עם רווח - בגרשיים בודדים)
FLEET_SH = "'צי הרחפנים'"
EVO_SH = "'תכולת ערכת EVO'"
AVATA_SH = "'תכולת ערכת AVATA'"
SOL_SH = "חיילים"
FROW0, FROW1 = 3, 52          # שורות נתוני הצי
LROW0, LROW1 = 3, 102         # שורות יומן
SROW0, SROW1 = 3, 62          # שורות חיילים
TYPES_TBL = "מקרא!$A$13:$C$14"

# ================================================================ סקירה
dash = wb.active
dash.title = "סקירה"
rtl(dash)
title(dash, "מערכת ניהול רחפנים – ארמו\"ן גדוד 28", 6)

def kpi(r, c, label, formula, fill=None):
    lc = dash.cell(r, c, label); lc.font = bold; lc.alignment = right; lc.border = border
    vc = dash.cell(r, c + 1, formula); vc.font = Font(bold=True, size=12)
    vc.alignment = center; vc.border = border
    if fill:
        vc.fill = fill
    return vc

for c, w in zip(range(1, 7), (18, 8, 4, 14, 10, 10)):
    dash.column_dimensions[L(c)].width = w

dash.cell(3, 1, "סיכום צי").font = Font(bold=True, size=13)
kpi(4, 1, "סה\"כ רחפנים", f"=COUNTA({FLEET_SH}!$B${FROW0}:$B${FROW1})", gray)
kpi(5, 1, "EVO", f'=COUNTIF({FLEET_SH}!$C${FROW0}:$C${FROW1},"EVO")')
kpi(6, 1, "AVATA", f'=COUNTIF({FLEET_SH}!$C${FROW0}:$C${FROW1},"AVATA")')

dash.cell(8, 1, "לפי סטטוס").font = Font(bold=True, size=13)
for i, s in enumerate(STATUSES):
    vc = kpi(9 + i, 1, s,
             f'=COUNTIF({FLEET_SH}!$E${FROW0}:$E${FROW1},"{s}")')
    dash.cell(9 + i, 1).fill = PatternFill("solid", fgColor=STATUS_FILL[s])

dash.cell(14, 1, "התראות").font = Font(bold=True, size=13)
alert1 = kpi(15, 1, "רחפנים תקולים",
             f'=COUNTIF({FLEET_SH}!$I${FROW0}:$I${FROW1},"תקול")')
alert2 = kpi(16, 1, "השאלות באיחור",
             f'=SUMPRODUCT(({FLEET_SH}!$E${FROW0}:$E${FROW1}="הושאל")*'
             f'({FLEET_SH}!$H${FROW0}:$H${FROW1}<>"")*'
             f'({FLEET_SH}!$H${FROW0}:$H${FROW1}<TODAY()))')
for cell in (alert1, alert2):
    dash.conditional_formatting.add(cell.coordinate,
        CellIsRule(operator="greaterThan", formula=["0"], fill=red, font=bold))

# מטריצה סוג x סטטוס
mrow = 4
dash.cell(mrow, 4, "סוג \\ סטטוס").fill = navy
dash.cell(mrow, 4).font = wbf; dash.cell(mrow, 4).alignment = center; dash.cell(mrow, 4).border = border
cols = ["EVO", "AVATA"]
for j, s in enumerate(STATUSES, start=5):
    c = dash.cell(mrow, j, s); c.fill = navy; c.font = wbf; c.alignment = center; c.border = border
    dash.column_dimensions[L(j)].width = 9
for i, t in enumerate(cols, start=1):
    rc = dash.cell(mrow + i, 4, t); rc.font = bold; rc.alignment = center; rc.border = border
    for j, s in enumerate(STATUSES, start=5):
        f = (f'=COUNTIFS({FLEET_SH}!$C${FROW0}:$C${FROW1},"{t}",'
             f'{FLEET_SH}!$E${FROW0}:$E${FROW1},"{s}")')
        cc = dash.cell(mrow + i, j, f); cc.alignment = center; cc.border = border
dash.cell(mrow + 3, 4, "סה\"כ").font = bold
dash.cell(mrow + 3, 4).alignment = center; dash.cell(mrow + 3, 4).border = border
for j in range(5, 9):
    col = L(j)
    cc = dash.cell(mrow + 3, j, f"=SUM({col}{mrow+1}:{col}{mrow+2})")
    cc.font = bold; cc.alignment = center; cc.border = border

dash.cell(18, 4, "עודכן:").font = bold
dash.cell(18, 5, "=TODAY()").number_format = "dd/mm/yyyy"

# ================================================================ צי הרחפנים
fleet = wb.create_sheet("צי הרחפנים")
rtl(fleet)
fleet_heads = ["מס׳", "מספר צ׳", "סוג", "מספר ערכה", "סטטוס", "מחזיק (חייל)",
               "תאריך השאלה", "תאריך החזרה יעד", "מצב", "פריטי צ׳ נדרשים", "הערות"]
title(fleet, "צי הרחפנים", len(fleet_heads))
head(fleet, 2, fleet_heads, [5, 13, 9, 11, 11, 16, 13, 15, 9, 13, 26])
for i, (serial, typ, kit, status, holder, d1, d2, cond, note) in enumerate(FLEET):
    r = FROW0 + i
    vals = [i + 1, serial, typ, kit, status, holder, D(d1), D(d2), cond,
            f'=IFERROR(VLOOKUP(C{r},{TYPES_TBL},3,FALSE),"")', note]
    for c, v in enumerate(vals, start=1):
        cell = fleet.cell(r, c, v); cell.alignment = center; cell.border = border
    fleet.cell(r, 7).number_format = "dd/mm/yyyy"
    fleet.cell(r, 8).number_format = "dd/mm/yyyy"
    fleet.cell(r, 2).alignment = right
# שורות ריקות עם מסגרת
for r in range(FROW0 + len(FLEET), FROW1 + 1):
    fleet.cell(r, 10, f'=IFERROR(VLOOKUP(C{r},{TYPES_TBL},3,FALSE),"")').alignment = center
    fleet.cell(r, 7).number_format = "dd/mm/yyyy"
    fleet.cell(r, 8).number_format = "dd/mm/yyyy"
    for c in range(1, 12):
        fleet.cell(r, c).border = border
fleet.freeze_panes = "A3"

# רשימות נפתחות
def add_dv(ws, formula, cells):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cells)

add_dv(fleet, '"EVO,AVATA"', f"C{FROW0}:C{FROW1}")
add_dv(fleet, '"במחסן,הושאל,בתיקון,מושבת"', f"E{FROW0}:E{FROW1}")
add_dv(fleet, '"תקין,תקול"', f"I{FROW0}:I{FROW1}")
add_dv(fleet, f"={SOL_SH}!$A${SROW0}:$A${SROW1}", f"F{FROW0}:F{FROW1}")
for c in ("G", "H"):
    fleet.column_dimensions[c].width = 14
# עיצוב מותנה לסטטוס/מצב/איחור
for s, color in STATUS_FILL.items():
    fleet.conditional_formatting.add(f"E{FROW0}:E{FROW1}",
        CellIsRule(operator="equal", formula=[f'"{s}"'],
                   fill=PatternFill("solid", fgColor=color)))
fleet.conditional_formatting.add(f"I{FROW0}:I{FROW1}",
    CellIsRule(operator="equal", formula=['"תקול"'], fill=red, font=bold))
fleet.conditional_formatting.add(f"I{FROW0}:I{FROW1}",
    CellIsRule(operator="equal", formula=['"תקין"'], fill=green))
fleet.conditional_formatting.add(f"H{FROW0}:H{FROW1}",
    FormulaRule(formula=[f'AND($E{FROW0}="הושאל",$H{FROW0}<>"",$H{FROW0}<TODAY())'],
                fill=red, font=bold))

# ================================================================ יומן השאלות
log = wb.create_sheet("יומן השאלות")
rtl(log)
log_heads = ["תאריך", "מספר צ׳", "סוג", "פעולה", "חייל", "מנפק", "הערות"]
title(log, "יומן השאלות והחזרות", len(log_heads))
head(log, 2, log_heads, [14, 13, 9, 12, 16, 16, 24])
for i, (d, serial, action, sol, issuer, note) in enumerate(LOG):
    r = LROW0 + i
    vals = [D(d), serial, f'=IFERROR(VLOOKUP(B{r},{FLEET_SH}!$B${FROW0}:$C${FROW1},2,FALSE),"")',
            action, sol, issuer, note]
    for c, v in enumerate(vals, start=1):
        cell = log.cell(r, c, v); cell.alignment = center; cell.border = border
    log.cell(r, 1).number_format = "dd/mm/yyyy"
for r in range(LROW0 + len(LOG), LROW1 + 1):
    log.cell(r, 3, f'=IFERROR(VLOOKUP(B{r},{FLEET_SH}!$B${FROW0}:$C${FROW1},2,FALSE),"")').alignment = center
    log.cell(r, 1).number_format = "dd/mm/yyyy"
    for c in range(1, 8):
        log.cell(r, c).border = border
log.freeze_panes = "A3"
add_dv(log, f"={FLEET_SH}!$B${FROW0}:$B${FROW1}", f"B{LROW0}:B{LROW1}")
add_dv(log, '"השאלה,החזרה"', f"D{LROW0}:D{LROW1}")
add_dv(log, f"={SOL_SH}!$A${SROW0}:$A${SROW1}", f"E{LROW0}:E{LROW1}")

# ================================================================ תכולת ערכה
def build_kit(name, items):
    ws = wb.create_sheet(name)
    rtl(ws)
    heads = ["#", "שם הפריט", "כמות", "פריט צ׳", "קיים?"]
    title(ws, name, len(heads))
    head(ws, 2, heads, [5, 52, 8, 9, 9])
    for i, (item, tz) in enumerate(items, start=1):
        r = 2 + i
        ws.cell(r, 1, i).alignment = center
        it = ws.cell(r, 2, item); it.alignment = right
        ws.cell(r, 3, 1).alignment = center
        tzc = ws.cell(r, 4, tz); tzc.alignment = center
        ws.cell(r, 5).alignment = center
        for c in range(1, 6):
            ws.cell(r, c).border = border
    last = 2 + len(items)
    add_dv(ws, '"כן,לא"', f"E3:E{last}")
    ws.conditional_formatting.add(f"E3:E{last}",
        CellIsRule(operator="equal", formula=['"כן"'], fill=green))
    ws.conditional_formatting.add(f"E3:E{last}",
        CellIsRule(operator="equal", formula=['"לא"'], fill=red, font=bold))
    ws.conditional_formatting.add(f"D3:D{last}",
        CellIsRule(operator="equal", formula=['"כן"'], fill=yellow))
    # סיכום
    s = last + 1
    ws.cell(s, 2, "סה\"כ פריטים:").font = bold; ws.cell(s, 2).alignment = right
    ws.cell(s, 3, f"=COUNTA(B3:B{last})").font = bold; ws.cell(s, 3).alignment = center
    ws.cell(s + 1, 2, "מתוכם פריטי צ׳:").font = bold; ws.cell(s + 1, 2).alignment = right
    ws.cell(s + 1, 3, f'=COUNTIF(D3:D{last},"כן")').font = bold; ws.cell(s + 1, 3).alignment = center
    ws.cell(s + 2, 2, "פריטים חסרים:").font = bold; ws.cell(s + 2, 2).alignment = right
    mc = ws.cell(s + 2, 3, f'=COUNTIF(E3:E{last},"לא")'); mc.font = bold; mc.alignment = center
    ws.conditional_formatting.add(mc.coordinate,
        CellIsRule(operator="greaterThan", formula=["0"], fill=red, font=bold))
    ws.freeze_panes = "A3"
    return last

evo_last = build_kit("תכולת ערכת EVO", EVO_ITEMS)
avata_last = build_kit("תכולת ערכת AVATA", AVATA_ITEMS)

# ================================================================ חיילים
sol = wb.create_sheet("חיילים")
rtl(sol)
title(sol, "חיילים", 3)
head(sol, 2, ["שם", "פלאפון", "הערות"], [22, 16, 26])
for i, name in enumerate(SOLDIERS):
    r = SROW0 + i
    sol.cell(r, 1, name).alignment = right
    for c in range(1, 4):
        sol.cell(r, c).border = border
for r in range(SROW0 + len(SOLDIERS), SROW1 + 1):
    for c in range(1, 4):
        sol.cell(r, c).border = border
sol.freeze_panes = "A3"

# ================================================================ מקרא
leg = wb.create_sheet("מקרא")
rtl(leg)
title(leg, "מקרא והגדרות", 4)
leg.column_dimensions["A"].width = 14
leg.column_dimensions["B"].width = 18
leg.column_dimensions["C"].width = 14
leg.cell(3, 1, "סטטוס רחפן:").font = bold
for i, s in enumerate(STATUSES):
    sw = leg.cell(4 + i, 1); sw.fill = PatternFill("solid", fgColor=STATUS_FILL[s]); sw.border = border
    leg.cell(4 + i, 2, s)
leg.cell(9, 1, "מצב:").font = bold
leg.cell(10, 1).fill = green; leg.cell(10, 1).border = border; leg.cell(10, 2, "תקין")
leg.cell(11, 1).fill = red; leg.cell(11, 1).border = border; leg.cell(11, 2, "תקול")
# טבלת סוגים (A13:C14) - נספרת מהתכולות, מוזנת ל-VLOOKUP בצי
for c, lab in enumerate(["סוג", "סה\"כ פריטים", "פריטי צ׳"], start=1):
    cc = leg.cell(12, c, lab); cc.fill = navy; cc.font = wbf; cc.alignment = center; cc.border = border
leg.cell(13, 1, "EVO").alignment = center
leg.cell(13, 2, f"=COUNTA({EVO_SH}!$B$3:$B${evo_last})").alignment = center
leg.cell(13, 3, f'=COUNTIF({EVO_SH}!$D$3:$D${evo_last},"כן")').alignment = center
leg.cell(14, 1, "AVATA").alignment = center
leg.cell(14, 2, f"=COUNTA({AVATA_SH}!$B$3:$B${avata_last})").alignment = center
leg.cell(14, 3, f'=COUNTIF({AVATA_SH}!$D$3:$D${avata_last},"כן")').alignment = center
for r in (13, 14):
    for c in range(1, 4):
        leg.cell(r, c).border = border
leg.cell(16, 1, "הוראות:").font = bold
notes = [
    "• כל רחפן = שורה בגיליון 'צי הרחפנים'.",
    "• בעת השאלה: עדכן סטטוס='הושאל', מחזיק ותאריכים, והוסף שורה ל'יומן השאלות'.",
    "• 'פריטי צ׳ נדרשים' מתעדכן אוטומטית לפי הסוג (מטבלת הסוגים כאן).",
    "• הדשבורד ('סקירה') מתעדכן לבד מהצי.",
    "• לספירת ערכה: סמן 'קיים?' בגיליון התכולה - החוסרים נספרים אוטומטית.",
]
for i, t in enumerate(notes):
    leg.cell(17 + i, 1, t).alignment = right
    leg.merge_cells(start_row=17 + i, start_column=1, end_row=17 + i, end_column=4)

OUT = "drone_management.xlsx"
wb.save(OUT)
print("נשמר:", OUT, "| גיליונות:", wb.sheetnames)
